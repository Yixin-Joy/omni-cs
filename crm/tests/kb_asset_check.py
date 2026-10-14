#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""知識庫「圖文」自檢：PDF 行級解析 + 圖片抽取/題注/綁定 + 命中帶圖 + 回覆只建一個任務。

涵蓋三件事：
  ① PDF 改走**行級解析**（pdfplumber）：頁碼/頁眉頁腳會被丟掉、
     **跨頁只是物理隔斷**（不產生段落分隔）。
  ② 圖片抽取（`ParsedAsset`）＋題注優先＋裝飾圖過濾。
  ③ 圖**必須綁到知識塊**，命中時一起帶出來（圖文結合輸出）。

★ CRM 側的三個差異（這支自檢也跟著改）：
  · 知識庫掛在商家底下 → 自檢建 `__CK__圖文自檢` 商家，跑完連商家一起清掉。
  · **`/console/kb/*` 尚未掛載進 CRM**（階段四決策：`kb_console.py` 仍是舊 API，
    接線等於整支重寫）。所以端點那幾段改成**直接呼叫 `kb_console` 的端點函式**
    （它們的參數就是 plain dict / int，不需要 request context）—— 驗到的還是同一份
    邏輯本體，等真的接線後把這幾行換成 HTTP 即可。
  · **送出搬到 Connector 了**：CRM 只負責「建一個外送任務」，所以 [7][8] 驗的是
    「命中帶圖 → 只建一個任務、內容是 text + assets」，而不是舊版的「送幾條訊息」。

⚠️ 合成 PDF 用 **英文語料**：內建的 Type1 Helvetica 沒有中文字型，
   餵中文會被抽成 `(cid:xxx)` 亂碼（中文的真實案例請看 `kb_check.py` 的 docx）。

可重複執行（自建臨時知識庫，跑完刪乾淨）。跑法（cwd 隨意）：
    PYTHONPATH='<root>\\.venv\\Lib\\site-packages' PYTHONUTF8=1 PYTHONIOENCODING=utf-8 \\
      <底層 python> crm/tests/kb_asset_check.py
"""
import asyncio
import json
import os
import sys

from _ck import check, report, sandbox_dir, ensure_business, purge_business  # noqa: E402

from app.rag import kb_chunk
from app.rag import kb_console
from app.rag import kb_db
from app.rag import kb_ingest
from app.rag import kb_parse
from app.rag import kb_search
from app.rag import rag_config

TMP = str(sandbox_dir())
# ⚠️ 名字要跟「使用者自己在用的庫」分開 —— 這個庫會被建立／刪除。
CHECK_COLLECTION = "圖文自檢庫(測試)"
BIZ = ensure_business("圖文自檢", "kb_asset_check 專用")


def body(res):
    """`_ok()` / `_err()` 回的是 Starlette Response → 自己解成 dict。"""
    raw = getattr(res, "body", None) or b"{}"
    try:
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return {}


# ★ 塊C（2026-10-14）：kb_console 端點已加登入身分 + 商家校驗。
#   自檢是**直接呼叫端點函式**（不走 HTTP，所以 FastAPI 的 Depends 不會生效），
#   必須自己補上 `sess` 與 `business_id`。這裡統一在模組層包一層，
#   呼叫點一行都不用改。
SUPER_SESS = {"active_role": "agent", "roles": ["admin", "agent"],
              "staff_id": 1, "is_super_agent": True,
              # ★ 寫入類端點（建/刪庫、改塊、上傳…）要通過 `_kb_edit_guard`
              "is_kb_editor": True}


def _wrap_endpoints(mod, biz):
    import functools
    import inspect
    for name in dir(mod):
        if name.startswith("_"):
            continue
        fn = getattr(mod, name)
        if not asyncio.iscoroutinefunction(fn):
            continue
        sig = inspect.signature(fn)
        if "sess" not in sig.parameters:
            continue

        @functools.wraps(fn)
        async def inner(*a, __fn=fn, __sig=sig, **kw):
            kw.setdefault("sess", SUPER_SESS)
            names = list(__sig.parameters)
            if "business_id" in names and kw.get("business_id") is None:
                # ⚠️ 呼叫端可能用**位置參數**傳到 business_id 那一格 → 別重複注入
                if len(a) <= names.index("business_id"):
                    kw["business_id"] = biz
            return await __fn(*a, **kw)

        setattr(mod, name, inner)


_wrap_endpoints(kb_console, BIZ)


# ---------- 造檔 ----------

def make_pdf_pages(path, pages, w=595, h=842):
    """pages = [[(y, text), ...], ...]；y 是**從頁面底部算**的基線座標。

    ⇒ 要讓某行落在「頁底」（跨頁續接的判斷依據），y 要小（例如 95）；
      要落在「頁首」，y 要大（例如 780）。
    """
    def esc(s):
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    n = len(pages)
    objs, kids = {}, []
    for i, lines in enumerate(pages):
        pno, cno = 4 + 2 * i, 5 + 2 * i
        kids.append(f"{pno} 0 R")
        content = "BT /F1 12 Tf\n"
        for (y, text) in lines:
            content += f"1 0 0 1 50 {y} Tm ({esc(text)}) Tj\n"
        content += "ET"
        objs[pno] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {w} {h}] "
                     f"/Resources << /Font << /F1 3 0 R >> >> /Contents {cno} 0 R >>")
        objs[cno] = f"<< /Length {len(content)} >>\nstream\n{content}\nendstream"
    objs[1] = "<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {n} >>"
    objs[3] = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    out = b"%PDF-1.4\n"
    offs = {}
    for i in sorted(objs):
        offs[i] = len(out)
        out += f"{i} 0 obj\n{objs[i]}\nendobj\n".encode()
    total = max(objs) + 1
    x = len(out)
    out += f"xref\n0 {total}\n0000000000 65535 f \n".encode()
    for i in range(1, total):
        out += (f"{offs[i]:010d} 00000 n \n" if i in offs else "0000000000 65535 f \n").encode()
    out += f"trailer\n<< /Size {total} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)


def make_docx_with_image(path, big=(800, 500), small=(40, 30), caption="figure 1-1 setup page"):
    from PIL import Image
    import docx
    from docx.shared import Inches

    p1 = os.path.join(TMP, "_img_big.png")
    p2 = os.path.join(TMP, "_img_small.png")
    Image.new("RGB", big, (220, 235, 250)).save(p1)
    Image.new("RGB", small, (255, 255, 255)).save(p2)

    d = docx.Document()
    d.add_heading("操作說明", level=1)
    d.add_paragraph("第一步：開啟設定頁面。")
    d.add_picture(p1, width=Inches(4))
    d.add_paragraph(caption)                 # ← 題注
    d.add_paragraph("第二步：輸入連線資訊。")
    d.add_picture(p2)                        # ← 太小 → 應被過濾
    d.add_paragraph("第三步：儲存。")
    d.save(path)
    return path


# ======================================================================
# [1] PDF 行級解析
# ======================================================================

print("=" * 74)
print(" 知識庫圖文自檢（PDF 行級解析 + 圖片抽取/綁定）")
print("=" * 74)

print("\n[1] PDF 行級解析：頁碼／頁眉頁腳過濾 + 跨頁接續")
pdf_path = os.path.join(TMP, "linelevel.pdf")
make_pdf_pages(pdf_path, [
    # p1：標頭 + 一句話，**最後一行在頁底**（y 小）
    [(780, "Product manual introduction."),
     (760, "The A100 headset costs 1280 and ships in"),
     (95, "two business days from our warehouse")],
    # p2：**第一行在頁首**（y 大）→ 應該跟 p1 最後一行接起來
    [(780, "and is covered by a 12 month warranty."),
     (760, "Contact support for details."),
     (95, "- 2 -")],                                    # ← 頁碼（頁腳帶 + 頁碼樣式）
    [(780, "Spare parts can be ordered separately."),
     (95, "- 3 -")],
])
parsed = kb_parse.parse(pdf_path, "pdf")

check(parsed.meta.get("pages") == 3, "讀到 3 頁", str(parsed.meta.get("pages")))
check(parsed.meta.get("dropped_header_footer") == 2,
      "兩個頁碼都被辨識並丟掉",
      f"dropped={parsed.meta.get('dropped_header_footer')}")
check(not any(("- 2 -" in b.text) or ("- 3 -" in b.text) or b.text.strip().isdigit()
              for b in parsed.blocks),
      "頁碼沒有混進任何知識塊",
      str([b.text for b in parsed.blocks if b.text.strip().isdigit()]))
check(parsed.meta.get("cross_page_joins") == 1,
      "跨頁續接被辨識（p1 尾 → p2 首 只有 1 處）",
      f"joins={parsed.meta.get('cross_page_joins')}")
joined = [b.text for b in parsed.blocks if "warehouse and is covered" in b.text]
check(bool(joined), "跨頁那句被接成同一段（不是兩個段落）",
      str(joined)[:90])
check(joined and joined[0].endswith("warranty."),
      "接起來之後整句完整到句尾")
check(all(b.page_no is not None for b in parsed.blocks),
      "每個塊都有 page_no（圖片才綁得上）",
      str([b.page_no for b in parsed.blocks]))
check(any(b.text.endswith(".") for b in parsed.blocks),
      "真正以句點結束的句子仍然自成一段（沒被過度合併）")
# 頁碼被丟掉，但**內文一字不少**
_all = "".join(b.text for b in parsed.blocks)
for needle in ("Product manual introduction.", "The A100 headset costs 1280",
               "Contact support for details.", "Spare parts can be ordered separately."):
    if needle not in _all:
        check(False, f"內文沒有遺漏：{needle[:30]}")

print("\n[2] 圖片抽取 + 題注優先 + 裝飾圖過濾")
docx_path = make_docx_with_image(os.path.join(TMP, "with_image.docx"))
pd = kb_parse.parse(docx_path, "docx")
check(len(pd.assets) == 2, f"抽出 2 張圖", f"實際 {len(pd.assets)}")
big = next((a for a in pd.assets if a.width >= 100), None)
small = next((a for a in pd.assets if a.width < 100), None)
check(big is not None and big.caption == "figure 1-1 setup page",
      "題注被當成預設描述（不必再花 AI）",
      repr(big.caption if big else None))
check(big is not None and not big.filtered_reason, "正常的圖沒有被過濾掉")
check(small is not None and bool(small.filtered_reason),
      "太小的圖被標成裝飾（仍保留記錄）", repr(small.filtered_reason if small else None))
check(all(a.anchor_block is not None for a in pd.assets),
      "每張圖都有綁定錨點（anchor_block）",
      str([a.anchor_block for a in pd.assets]))

# ======================================================================
# [3] 入庫：落地 + 去重 + 綁定分塊
# ======================================================================

print("\n[3] 入庫：圖片落地 + 綁定知識塊")
cid = None
for c in kb_db.list_collections(business_id=BIZ):
    if c["name"] == CHECK_COLLECTION:
        cid = c["id"]
if not cid:
    cid = kb_db.create_collection(CHECK_COLLECTION, "圖文自檢用，可刪",
                                  business_id=BIZ)["id"]
else:
    kb_db.update_collection(cid, {"enabled": True})     # 自檢庫一定要啟用

with open(docx_path, "rb") as f:
    up = kb_ingest.save_upload("with_image.docx", f.read(), "", cid)
did = up["doc"]["id"]
kb_ingest.parse_and_chunk(did)

assets = kb_db.list_assets(doc_id=did)
check(len(assets) == 2, f"入庫後 kb_asset 有 2 筆", f"實際 {len(assets)}")
check(all(a["rel_path"] for a in assets), "每筆都有落地路徑")
on_disk = [a for a in assets
           if os.path.exists(os.path.join(rag_config.RAG_MEDIA_DIR,
                                          os.path.basename(a["rel_path"])))]
check(len(on_disk) == len(assets), "圖檔真的寫到公開目錄 media/kb_assets/ 底下",
      f"{len(on_disk)}/{len(assets)}")
check(all(a["doc_id"] == did for a in assets), "圖掛在這份文件底下")
check(any(a["caption_source"] == "doc" for a in assets),
      "題注來源標成 doc")
check(any(not a["enabled"] and a["filtered_reason"] for a in assets),
      "被過濾的圖 enabled=false（不會被帶出去）")
check(all((a["chunk_count"] or 0) > 0 for a in assets if a["enabled"]),
      "**啟用中的圖都綁到了知識塊**（圖文結合輸出的前提）",
      str([(a["id"], a["chunk_count"]) for a in assets]))

# 重複入庫不該產生重複的圖（sha256 去重）
kb_ingest.parse_and_chunk(did)
check(len(kb_db.list_assets(doc_id=did)) == 2,
      "重新分塊不會產生重複圖片（sha256 去重）",
      f"{len(kb_db.list_assets(doc_id=did))} 筆")

# ======================================================================
# [4] 圖片端點（直接呼叫處理函式，因為 /console/kb/* 尚未掛載）
# ======================================================================

print("\n[4] 圖片端點函式（列表 / 人工描述 / AI 不覆蓋人工 / 停用不帶出）")

r = asyncio.run(kb_console.list_assets(doc_id=str(did), include_filtered="1"))
d = body(r)
check(len(d.get("assets", [])) == 2, "list_assets 回 2 張", f"實得 {len(d.get('assets', []))}")
check(d.get("stats", {}).get("with_caption", 0) >= 1, "統計：有描述的張數 > 0",
      str(d.get("stats")))
d2 = body(asyncio.run(kb_console.list_assets(doc_id=str(did), include_filtered="0")))
check(len(d2.get("assets", [])) == 1, "include_filtered=0 只回內容圖")

aid = next(a["id"] for a in assets if a["caption_source"] == "doc")
rf = asyncio.run(kb_console.asset_file(aid))
# ⚠️ FileResponse 是**串流**的：`res.body` 本來就是空的，要看它指向的檔案。
_rf_path = getattr(rf, "path", "")
check(getattr(rf, "status_code", 200) == 200 and _rf_path
      and os.path.getsize(_rf_path) > 100,
      "asset_file() 回得到圖檔（FileResponse 指向真實檔案）",
      f"{os.path.getsize(_rf_path) if _rf_path else 0} bytes")
check(getattr(asyncio.run(kb_console.asset_file(999999)), "status_code", 200) == 404,
      "不存在的圖回 404")

r = asyncio.run(kb_console.set_asset_caption({"id": aid, "caption": "人工寫的描述"}))
check(body(r)["asset"]["caption_source"] == "manual",
      "人工描述寫入，來源標 manual")

# ⚠️ 這條**不會真的呼叫模型**：端點在 manual 時直接跳過 → 可以放心跑
rr = (body(asyncio.run(kb_console.ai_asset_caption({"id": aid}))).get("results")
      or [{}])[0]
check(rr.get("skipped") and not rr.get("ok"),
      "AI 描述**不會覆蓋**人工寫的描述（要覆蓋得帶 force=1）",
      str(rr.get("error"))[:50])
check(kb_db.get_asset(aid)["caption"] == "人工寫的描述", "人工描述沒有被動到")

r = asyncio.run(kb_console.set_asset_caption({"id": aid, "caption": ""}))
check(body(r)["asset"]["caption_status"] == "pending", "清空描述 → 回到待填")

# 停用後不該被檢索帶出
small_id = next(a["id"] for a in assets if not a["enabled"])
bound = kb_db.assets_for_chunks([c["id"] for c in kb_db.list_chunks(did)])
check(small_id not in [x["id"] for arr in bound.values() for x in arr],
      "被停用的圖不會出現在 assets_for_chunks")

# ★ 記錄現況：端點還沒掛進 main.app，所以外部（含工作台）打不到。
import main as crm_main      # noqa: E402
_kb_routes = [getattr(rt, "path", "") for rt in crm_main.app.routes
              if str(getattr(rt, "path", "")).startswith("/console/kb")]
if not _kb_routes:
    print("  [未接線] main.app 目前沒有 /console/kb/* —— 上面這幾條是直接呼叫處理函式，")
    print("           所以驗到邏輯本體；真的接線後請把這幾行改成打 HTTP。")

# ======================================================================
# [5] 命中帶圖
# ======================================================================

print("\n[5] 檢索命中 → 連圖一起帶出來")
res_embed = kb_ingest.embed_document(did, "pending")
kb_ingest.activate_document(did)
kb_search.invalidate()
check(res_embed.get("done", 0) > 0, "向量化完成", str(res_embed))

sr = kb_search.search("設定頁面怎麼開啟", business_id=BIZ)
check(sr.passed, "檢索命中", f"top1={sr.top1:.3f}")
check(bool(sr.assets), "SearchResult 帶出了圖片", f"{len(sr.assets)} 張")
check(all("url" in a for a in sr.assets), "每張圖都有取圖 url")
hit_with_assets = [h for h in sr.hits if h.assets]
check(bool(hit_with_assets), "至少有一個入選塊帶著圖（Hit.assets）")
js = sr.hits[0].as_dict() if sr.hits else {}
check("assets" in js, "Hit.as_dict() 含 assets（寫進 matched_rows，紀錄頁才看得到）")

# ======================================================================
# [6] 圖片內嵌 markdown（2026-10-12）：圖片 → chunk 文字 / 只對文字向量化
# ======================================================================

print("\n[6] 圖片內嵌 markdown + 圖片庫可被檢索")
from app.rag import kb_chunk as _kc          # noqa: E402
from app.rag import rag_config as _rc        # noqa: E402

kb_ingest.embed_document(did, "pending")
kb_ingest.activate_document(did)
kb_search.invalidate()

chunks = kb_db.list_chunks(did)
withmd = [c for c in chunks if "![" in (c["text"] or "")]
check(bool(withmd), "chunk 文字裡有內嵌的 markdown 圖片標記",
      (withmd[0]["text"][-90:] if withmd else "（沒有）"))
check(any("/media/" in (c["text"] or "") for c in chunks),
      "標記用的是公開路徑 /media/…（前端直接渲染）")
check(any("figure 1-1 setup page" in (c["text"] or "") for c in chunks),
      "標記的 alt 用了題注文字（沒題注才留空）")

check(_kc.strip_md_images("前\n![圖 1](/media/a.png)\n後") == "前\n\n後",
      "strip_md_images 剝掉圖片標記只留文字（向量化只吃文字）")
check(_kc.strip_md_images("沒有圖片的文字") == "沒有圖片的文字", "沒有標記時原樣回傳")

a_on = [a for a in kb_db.list_assets(doc_id=did) if a["enabled"]]
check(bool(a_on) and all(_rc.media_url_of_rel(a["rel_path"]).startswith("/media/")
                         for a in a_on),
      "啟用中的圖都有公開網址 /media/…",
      [_rc.media_url_of_rel(a["rel_path"]) for a in a_on])
check(all(os.path.exists(os.path.join(_rc.RAG_MEDIA_DIR, os.path.basename(a["rel_path"])))
          for a in a_on),
      "圖檔真的落在公開目錄 media/kb_assets/（不是整包 media 都公開）")

imgcol = kb_db.get_image_collection(cid)
check(bool(imgcol) and imgcol["kind"] == "image" and imgcol["image_of"] == cid,
      "文檔庫有對應的圖片庫（kind=image / image_of=cid）",
      (f"#{imgcol['id']} {imgcol['name']}" if imgcol else "（沒有）"))
check(bool(imgcol) and all(a["collection_id"] == imgcol["id"] for a in a_on),
      "從文件抽出的圖落在該文檔庫的圖片庫",
      [(a["id"], a["collection_id"]) for a in a_on])

# 寫描述（模擬真實流程：描述寫好 → 立刻向量化 → 之後就查得到）
own_aid = a_on[0]["id"] if a_on else None
r = asyncio.run(kb_console.set_asset_caption(
    {"id": own_aid, "caption": "figure 1-1 setup page"}))
check(getattr(r, "status_code", 200) == 200
      and body(r)["asset"]["caption_source"] == "manual",
      "圖片描述寫入成功（來源 manual）")
check(kb_db.get_asset(own_aid)["embedding_status"] == "done",
      "★ 寫完描述**立刻向量化**（圖片能被檢索的前提；以前完全沒這一步）")
kb_search.invalidate()

r = kb_search.search("figure 1-1 setup page", business_id=BIZ)
imhits = [h for h in r.candidates if h.kind == "image"]
check(bool(imhits), "★ 圖片庫條目進得了候選池（描述可被檢索）— 以前完全搜不到",
      [(h.kind, round(h.final, 3)) for h in r.candidates[:3]])
own = [h for h in imhits if h.asset_id == own_aid]
check(bool(own), "這張圖的條目在候選池裡（text 就是 markdown，前端渲染即可）",
      (own[0].text if own else [h.text[:44] for h in imhits[:3]]))
check(bool(own) and own[0].text.startswith("![figure 1-1"),
      "圖片命中的 text 的 alt 用了描述", (own[0].text[:60] if own else ""))

# 去重：綁在入選塊上的圖，不該同時以「圖片條目」的身分再出現一次
bound_ids = set()
for _arr in (kb_db.assets_for_chunks([h.chunk_id for h in r.hits
                                      if h.kind != "image"]).values() or []):
    bound_ids |= {int(x["id"]) for x in _arr}
dup = [h for h in r.hits if h.kind == "image" and h.asset_id in bound_ids]
check(not (bound_ids and dup),
      "同一張圖同時被 chunk 與圖片庫命中時只出一次",
      f"綁定中的圖 {sorted(bound_ids)} / 重複出現的圖片命中 {[h.asset_id for h in dup]}")


kb_db.set_collection_enabled(cid, False)
kb_search.invalidate()
r2 = kb_search.search("figure 1-1 setup page", business_id=BIZ)
check(not [h for h in r2.candidates if h.kind == "image" and h.asset_id == own_aid],
      "文檔庫停用後它的圖片庫也檢索不到（啟停用連動）")
check(kb_db.get_collection(imgcol["id"])["enabled"] is False, "圖片庫 enabled 被連動關閉")
kb_db.set_collection_enabled(cid, True)
kb_search.invalidate()
check(bool([h for h in kb_search.search("figure 1-1 setup page", business_id=BIZ).candidates
            if h.kind == "image" and h.asset_id == own_aid]), "恢復啟用後又能檢索到")

# ---------- [7] 送出端的策略判斷（純函式） ----------
#
# 舊單體版的症狀（小jo 實測）：知識庫命中帶圖時，客戶端收到
#   「一條不帶圖的純文字」+「一條一模一樣、但帶圖的」。
# 修法：能帶 caption 的渠道（WhatsApp）把文字塞進圖片的 caption 一起發，只發一條。
#
# ★ CRM 側：**真正發送搬到 Connector 了** —— CRM 只負責「建一個外送任務」
#   （內容 = text + assets），所以這裡驗的是 CRM 端的決策
#   （`assets_inline_caption()` 純函式）＋ [8] 的「只建一個任務」。
print("\n[7] 送出端策略：文字能不能塞進圖片 caption（純函式）")

from app.autoreply import autoreply          # noqa: E402

_a = kb_db.get_asset(own_aid)
_asset = dict(_a or {})
_text = "伺服器端轉發成功的日誌如下圖所示。"

check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset], _text) is True,
      "WhatsApp 私訊 + 有圖 + 有文字 → 走「圖文一條」")
check(autoreply.assets_inline_caption("wechat", "dm", [_asset], _text) is False,
      "微信圖片沒有 caption → 不宣稱能一條（平台限制，只能文字+圖）")
check(autoreply.assets_inline_caption("ig", "dm", [_asset], _text) is False,
      "IG 的 attachment 沒有 caption → 不宣稱能一條")
check(autoreply.assets_inline_caption("fb", "dm", [_asset], _text) is False,
      "FB 同上（attachment 沒有 caption）")
check(autoreply.assets_inline_caption("whatsapp", "comment", [_asset], _text) is False,
      "貼文留言不發圖（公開場合）")
check(autoreply.assets_inline_caption("whatsapp", "dm", [], _text) is False,
      "沒有圖 → 沒什麼好 inline 的")
check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset], "   ") is False,
      "只有空白文字 → 沒有東西可以當 caption")
check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset],
                                      "字" * (autoreply.CAPTION_MAX + 1)) is False,
      f"文字超過 caption 上限（{autoreply.CAPTION_MAX}）→ 不硬塞")
check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset],
                                      "字" * autoreply.CAPTION_MAX) is True,
      "剛好等於上限 → 仍算一條")

# 任務內容的組裝（`_asset_items` / `_strip_images` / `_assets_caption`）
_items = autoreply._asset_items([_asset, {"rel_path": ""}, None])
check(len(_items) == 1 and _items[0]["rel_path"] == _asset["rel_path"],
      "組任務內容時濾掉沒有 rel_path 的空項")
check(autoreply._strip_images(_text + "\n![圖 1](/media/x.png)") == _text,
      "★ 任務裡的文字已剝掉 markdown 圖片標記（客戶端不認 markdown）")
check(autoreply._strip_images("![圖 1](/media/x.png)") == "",
      "純圖片回覆 → 剝完是空的（呼叫端改用題注當說明）")
check(autoreply._assets_caption([{"caption": "圖 1 截圖"}]) == "請參考下圖（圖 1 截圖）。",
      "純圖片回覆的文字說明用題注")
check(autoreply._assets_caption([]) == "請參考下圖。", "沒題注 → 用一句通用話")

# ---------- [8] CRM 端到端：命中帶圖 → 只建「一個」任務 ----------
#
# 把引擎與建任務都換成假的（**不真的發任何訊息、不真的寫任務**），
# 只看 handle_inbound 建了幾個任務、任務內容長什麼樣。
print("\n[8] 端到端：命中帶圖時只建一個任務（不是文字任務 + 圖片任務）")

from app.llm import reply_engine          # noqa: E402

_ANSWER_TEXT = "伺服器端轉發成功的日誌如下圖所示。\n![圖 1 截圖](/media/placeholder.png)"
_enq: list[dict] = []
_logs: list[dict] = []


def _install_fakes():
    """把 handle_inbound 的外部依賴全部換成假的；回傳還原函式。"""
    from app.autoreply import autoreply_db

    saved = {("autoreply", n): getattr(autoreply, n)
             for n in ("is_enabled", "_guard", "_mark", "_log", "enqueue_reply")}
    saved[("reply_engine", "answer")] = reply_engine.answer
    for n in ("claim_message", "get_conversation_history", "resolve_scope_set"):
        saved[("autoreply_db", n)] = getattr(autoreply_db, n)

    async def _fake_answer(question, **kw):
        return reply_engine.Answer(reply_engine.TIER_RAG, text=_ANSWER_TEXT,
                                   assets=[dict(_asset)], detail="測試")

    def _fake_enqueue(**kw):
        _enq.append(kw)
        return {"task_id": 424242}

    async def _fake_history(*a, **kw):
        return []

    autoreply.is_enabled = lambda biz, ch, sc, acc=0: True
    autoreply._guard = lambda biz, end_user, ignore=False: ""
    autoreply._mark = lambda *a, **kw: None
    autoreply._log = lambda payload, claim_key=None: _logs.append(payload)
    autoreply.enqueue_reply = _fake_enqueue
    reply_engine.answer = _fake_answer
    autoreply_db.claim_message = lambda key, payload: True
    autoreply_db.get_conversation_history = _fake_history
    autoreply_db.resolve_scope_set = lambda biz, scope="dm", post_config_id=0: None

    def _restore():
        for (mod, name), val in saved.items():
            if mod == "autoreply":
                setattr(autoreply, name, val)
            elif mod == "reply_engine":
                setattr(reply_engine, name, val)
            else:
                setattr(autoreply_db, name, val)

    return _restore


def _run_inbound(channel):
    return asyncio.run(autoreply.handle_inbound(
        business_id=BIZ, channel=channel, scope="dm",
        text="服务器端转发成功日志长啥样？", to="85200000000",
        conversation_key="__kb_asset_check__", ignore_cooldown=True))


_restore = _install_fakes()
try:
    for _ch, _want_inline in (("whatsapp", True), ("wechat", False), ("ig", False)):
        _enq.clear()
        _logs.clear()
        _res = _run_inbound(_ch)
        check(_res.get("state") == "hit",
              f"{_ch}：handle_inbound 走命中分支", str(_res)[:70])
        check(len(_enq) == 1,
              f"★ {_ch}：只建 1 個任務（不是文字任務 + 圖片任務各一個）",
              f"建了 {len(_enq)} 個")
        _kw = _enq[0] if _enq else {}
        check(_kw.get("text") == "伺服器端轉發成功的日誌如下圖所示。",
              f"{_ch}：任務文字已剝掉 markdown 圖片標記", repr(_kw.get("text"))[:52])
        _as = _kw.get("assets") or []
        check(len(_as) == 1 and _as[0].get("rel_path") == _asset["rel_path"],
              f"{_ch}：圖以 assets 掛在**同一個**任務上")
        check(_kw.get("business_id") == BIZ and _kw.get("platform") == _ch
              and _kw.get("to") == "85200000000",
              f"{_ch}：任務帶對商家／平台／收件人")
        _detail = str((_logs[0] if _logs else {}).get("detail") or "")
        check(("文字隨圖" in _detail) is _want_inline,
              f"{_ch}：日誌{'標了' if _want_inline else '沒標'}「文字隨圖」",
              _detail[:60])
        check(str((_logs[0] if _logs else {}).get("send_status")) == "queued",
              f"{_ch}：命中記錄的 send_status = queued（真實結果由 Connector 回調）")
finally:
    _restore()

# ---------- [9] 標題識別 + 表格結構化（2026-10-12，「分塊效果不好」的根因） ----------
#
# 症狀（小jo 回報 doc#319）：格式清晰的 PDF，分塊卻一團糟 ——
#   · 4 個圖片題注被當成章節標題，把後面 51/74 個 block 的 heading_path 全汙染了
#   · 12 個真正的編號標題（「1.1 产品简介」）一個都沒認出來
#   · 標題與正文黏成一行（「1.1 产品简介本智能办公助手系统是面向…」）
#   · 表格被壓成流水文字（「故障现象 报错代码 核心原因 解决方案」後面接著糊掉的內容）
print("\n[9] 標題識別 + 表格結構化（分塊品質的關鍵）")

# 編號式章節標題：要認得出來，也要估對層級
for t, lv in [("1. 产品概述", 1), ("3.2 完整安装步骤（图文对应）", 2),
              ("1.1.2 细节说明", 3), ("5. 常见故障与解决方案（RAG 高频问答库）", 1)]:
    got = kb_parse._numbered_heading_level(t)
    check(got == lv, f"編號標題「{t}」→ 第 {lv} 層", f"實際 {got}")

# ★ 這些**不能**當標題：列表項（含句中標點）、版本號（數字後沒空白）
for t in ["1. 关闭电脑防火墙及杀毒软件拦截，避免安装包解压、组件注册被拦截",
          "2. 首次登录强制修改密码，密码需包含大小写字母、数字、8位及以上",
          "11.0及以上",
          "4. 初始化完成后自动进入主功能界面，可正常使用检索、问答功能"]:
    got = kb_parse._numbered_heading_level(t)
    check(got == 0, f"不算標題「{t[:20]}…」", f"實際 {got}")


def _ln(text, **kw):
    d = {"text": text, "x0": 90, "x1": 300, "top": 200, "bottom": 212}
    d.update(kw)
    return d


check(kb_parse._looks_heading(_ln("圖 2 文件解压完成目录示意图"), 595) is False,
      "★ 圖片題注不當標題（它由 _match_caption 綁到圖上）")
check(kb_parse._looks_heading(_ln("1.1 产品简介"), 595) is True,
      "★ 左對齊的編號標題要認得出來（原本只認置中 → 全漏）")
check(kb_parse._looks_heading(_ln("11.0及以上", x0=228, x1=284, in_table=True), 595) is False,
      "★ 表格內的欄位文字不當標題（偏右會像置中）")
check(kb_parse._looks_heading(_ln("第一章 总则"), 595) is True, "「第X章」仍認得")
check(kb_parse._heading_level_of("第三章 总则") == 1, "_heading_level_of 認得第X章")

# 表格 → markdown（列／欄關係要保住）
md = kb_parse._rows_to_markdown([["故障现象", "报错代码"], ["软件安装完成后无\n法启动", "E001"]])
check("| 故障现象 | 报错代码 |" in md, "表格轉成 markdown 表頭", repr(md.split("\n")[0]))
check("| 软件安装完成后无法启动 | E001 |" in md,
      "★ 儲存格內換行接回同一句（中文不加空格，否則「无法启动」檢索不到）",
      repr(md.split("\n")[-1]))
check(kb_parse._rows_to_markdown([["只有一欄"], ["x"]]) == "",
      "只有一欄的「表」判為誤判 → 不轉", "（免得把版心框當表格）")
check(kb_parse._rows_to_markdown([]) == "" and kb_parse._rows_to_markdown(None) == "",
      "空表格不炸")

# 前綴／內文去重（同一句別出現兩次）
check(kb_chunk._same_as_head("RAG问答测试专用｜图文版用户说明文档",
                             "RAG 问答测试专用｜图文版用户说明文档\n文档用途：…"),
      "前綴與內文開頭重複（忽略空白）→ 判為重複")
check(not kb_chunk._same_as_head("產品手冊", "耳機的保固期是 12 個月"),
      "不相干的前綴與內文 → 不判重複")
check(kb_chunk._drop_dup_heading("3.3 首次登录初始化设置\n软件首次打开需完成账号登录",
                                 "3. 软件安装 › 3.3 首次登录初始化设置")
      == "软件首次打开需完成账号登录",
      "★ 內文開頭的章節標題會被去掉（標題已在 prefix 裡）")
check(kb_chunk._drop_dup_heading("3.3 首次登录初始化设置", "3.3 首次登录初始化设置")
      == "3.3 首次登录初始化设置",
      "整塊就只有那個標題時 → 保留（去掉就空了）")

# ---------- 收尾 ----------
print("\n[10] 清理…")
kb_ingest.delete_document(did)
left = kb_db.asset_stats(doc_id=did).get("total") or 0
check(left == 0, "刪文件時圖片一起清掉", f"殘留 {left}")

# 自檢商家用完就清（庫＋它的圖片庫＋文件＋圖，一次清乾淨）
purge_business(BIZ)
left_cols = kb_db.list_collections(business_id=BIZ)
check(not left_cols, "自檢商家已被清空（含圖片庫／文件／圖檔紀錄）",
      f"剩 {len(left_cols)} 個庫")

raise SystemExit(report("知識庫圖文自檢"))
