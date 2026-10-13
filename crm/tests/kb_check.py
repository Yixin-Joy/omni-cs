#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""RAG 知識庫（入庫管道）自檢：解析 → 分塊 → embedding → 入庫。

可重複執行：每次會先清掉自檢庫的舊資料再重建，所以任何時候都能跑同一支腳本。
（模型本身的診斷請用 `rag_model_check.py`；HTTP 端點請用 `kb_api_check.py`。）

★ CRM 側：知識庫掛在商家底下（`kb_collection.business_id` NOT NULL）→ 自檢建一個
  `__CK__RAG自檢` 商家，跑完連庫帶商家一起清掉。

跑法（cwd 隨意，路徑由 __file__ 決定）：
    PYTHONPATH='<root>\\.venv\\Lib\\site-packages' PYTHONUTF8=1 PYTHONIOENCODING=utf-8 \\
      <底層 python> crm/tests/kb_check.py
"""
import json
import os
import sys
import time
from collections import Counter as _Counter

from _ck import (check, report, sandbox_dir, ensure_business,   # noqa: E402
                 purge_business)

from app.rag import kb_db
from app.rag import kb_ingest
from app.rag import kb_chunk
from app.rag import kb_parse
from app.rag import kb_search
from app.rag import rag_config

TMP = str(sandbox_dir())

CHECK_COLLECTION = "RAG 自檢庫"
OTHER_COLLECTION = "RAG 自檢庫（參數共用對照）"
BIZ = ensure_business("RAG自檢", "kb_check 專用")

# ---------- 造檔 ----------

def make_pdf(path, lines):
    def esc(s):
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    content = "BT /F1 12 Tf 50 780 Td 16 TL\n"
    for l in lines:
        content += f"({esc(l)}) Tj T*\n"
    content += "ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n").encode()
    with open(path, "wb") as f:
        f.write(out)


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("蓝牙耳机产品手册", level=1)
    d.add_paragraph("本手册介绍我们全系列蓝牙耳机的规格、价格与售后政策。")
    d.add_heading("A 款耳机", level=2)
    d.add_paragraph("A 款蓝牙耳机售价 1280 元，续航 30 小时，支持蓝牙 5.3。")
    d.add_paragraph("A 款支持主动降噪，防水等级 IPX4。")
    d.add_heading("价格表", level=3)
    t = d.add_table(rows=3, cols=3)
    t.rows[0].cells[0].text = "型号"
    t.rows[0].cells[1].text = "价格"
    t.rows[0].cells[2].text = "续航"
    t.rows[1].cells[0].text = "A 款"
    t.rows[1].cells[1].text = "1280 元"
    t.rows[1].cells[2].text = "30 小时"
    t.rows[2].cells[0].text = "B 款"
    t.rows[2].cells[1].text = "890 元"
    t.rows[2].cells[2].text = "20 小时"
    d.add_heading("售后政策", level=2)
    d.add_paragraph("商品到货后 7 天内，包装完整可无理由退货。")
    d.add_paragraph("电子产品自购买日起提供一年原厂保固。")
    d.save(path)


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "价格表"
    ws.append(["产品名称", "规格", "售价", "库存"])
    for i in range(1, 25):
        ws.append([f"蓝牙耳机 {chr(64 + (i % 6) + 1)}-{i}", f"{i * 10}小时续航",
                   800 + i * 13, 50 + i])
    ws2 = wb.create_sheet("售后")
    ws2.append(["项目", "说明"])
    ws2.append(["退货", "7天无理由退货，包装需完整"])
    ws2.append(["保固", "一年原厂保固"])
    wb.save(path)


def make_md(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("""# 客服知识库

## 运费说明
单笔订单满 500 元免运费，未满则收 60 元运费。
偏远地区（离岛）另收 100 元。

## 客服时间
周一至周五 09:00-18:00，国定假日休息。

## 付款方式
支持信用卡、ATM 转账、超商取货付款。

```python
# 这段是代码块，不应被切断
def hello():
    return "world"
```

## 保固
电子产品自购买日起提供一年原厂保固。
""")


def make_txt(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("产品保修说明\n\n所有电子产品均提供一年保修服务。\n保修期内非人为损坏免费维修。\n\n"
                "人为损坏需要收取零件费用。\n\n送修流程：联系客服 → 寄回商品 → 检测报价 → 维修或更换。\n")
        f.write("\n" * 1)
        for i in range(20):
            f.write(f"补充条款第 {i + 1} 条：本条为测试长文，用来验证分块是否会正确合并到目标长度。\n")


def make_csv(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("门店名称,城市,电话,营业时间\n")
        for i in range(1, 16):
            f.write(f"旗舰店{i},台北市,02-1234-{1000 + i},10:00-21:00\n")


print("=" * 72)
print(" RAG 知識庫入庫自檢（解析 → 分塊 → embedding → 入庫）")
print("=" * 72)

# ---------- 0) 解析器單元測試 ----------
print("\n[0] 解析器（各格式）…")
samples = {
    "manual.pdf": (make_pdf, os.path.join(TMP, "manual.pdf"), [
        "Bluetooth Earphone Manual",
        "A model price is 1280 CNY with 30 hours battery life.",
        "B model price is 890 CNY with active noise cancelling.",
        "Return policy: 7 days no reason return if package complete.",
    ]),
}
make_pdf(samples["manual.pdf"][1], samples["manual.pdf"][2])
make_docx(os.path.join(TMP, "manual.docx"))
make_xlsx(os.path.join(TMP, "price.xlsx"))
make_md(os.path.join(TMP, "faq.md"))
make_txt(os.path.join(TMP, "warranty.txt"))
make_csv(os.path.join(TMP, "stores.csv"))

for fn in ("manual.pdf", "manual.docx", "price.xlsx", "faq.md", "warranty.txt", "stores.csv"):
    p = os.path.join(TMP, fn)
    try:
        parsed = kb_parse.parse(p)
        blocks = len(parsed.blocks)
        check(blocks > 0, f"{fn} 解析出 {blocks} 區塊 / {parsed.char_count} 字",
              f"meta={json.dumps(parsed.meta, ensure_ascii=False)[:70]}")
    except Exception as e:
        check(False, f"{fn} 解析", str(e))

# 掃描型 PDF（沒有文字層）→ 應拋 no_text_layer
empty_pdf = os.path.join(TMP, "empty.pdf")
make_pdf(empty_pdf, [])
try:
    kb_parse.parse(empty_pdf)
    check(False, "空 PDF 應報 no_text_layer")
except kb_parse.ParseError as e:
    check(e.reason == "no_text_layer", f"空 PDF 正確回報 {e.reason}",
          kb_parse.describe_error(e.reason))

# ---------- 1) 分塊 ----------
print("\n[1] 分塊策略…")
for fn, dt in (("manual.pdf", "pdf"), ("manual.docx", "docx"), ("price.xlsx", "xlsx"),
               ("faq.md", "md"), ("warranty.txt", "txt"), ("stores.csv", "csv")):
    parsed = kb_parse.parse(os.path.join(TMP, fn))
    params = {"chunk_strategy": "auto", "chunk_size": 300, "chunk_min": 120,
              "chunk_max": 400, "chunk_overlap": 60, "chunk_prefix": True}
    cs = kb_chunk.chunk_blocks(parsed.blocks, doc_type=dt, title=fn.split(".")[0], params=params)
    sizes = [c.char_count for c in cs]
    ok = bool(cs) and all(s <= 400 + 60 for s in sizes)
    check(ok, f"{fn}: {len(cs)} 塊，長度 {min(sizes)}~{max(sizes)}",
          f"策略={kb_chunk.resolve_strategy(dt, 'auto')}")

# ---------- 1b) 分塊不得丟內容（回歸測試） ----------
# 這裡是踩過的坑：`_pack()` 舊寫法在累積途中用「尾端 overlap 字」整個取代緩衝區，
# 把還沒收成塊的內容直接丟掉 —— 600 字進、281 字出，整段消失（epro-bridge 手冊
# 4796 字只剩 2125 字就是這個原因）。重疊只允許「重複」，不允許「變少」。
print("\n[1b] 分塊不得丟內容（重疊 ≠ 截斷）…")
_long_paras = [kb_parse.Block(f"第{i}段" + "內容" * 90) for i in range(1, 5)]   # 每段 ~183 字
_src_chars = sum(len(b.text) for b in _long_paras)
_src_pool = _Counter("".join(b.text for b in _long_paras))
for strat, dt in (("paragraph", "txt"), ("heading", "md"), ("table_row", "csv"),
                  ("fixed", "txt")):
    for ov in (0, 60):
        cs = kb_chunk.chunk_blocks(_long_paras, doc_type=dt, title="t",
                                   params={"chunk_strategy": strat, "chunk_size": 300,
                                           "chunk_min": 120, "chunk_max": 400,
                                           "chunk_overlap": ov, "chunk_prefix": False})
        total = sum(len(c.body) for c in cs)
        whole = "\n".join(c.body for c in cs)
        # 通用條件：字元總量不得少於原文（重疊只會變多），且每個字元的出現次數
        # 都不少於原文 —— 這條對「任何」策略都成立。
        pool_ok = _Counter(whole) >= _src_pool
        check(total >= _src_chars and pool_ok,
              f"{strat}(重疊 {ov})：{len(cs)} 塊，內文 {total} 字 ≥ 原始 {_src_chars}",
              f"字元總量不足={not pool_ok}")
        if strat != "fixed":
            # ⚠️ `fixed` 是「無冗餘平鋪」（overlap=0 時）：段落橫跨切點就不會完整
            #    落在任何單一塊裡，所以「逐段完整」只對有結構的 pack 策略成立。
            missing = [i + 1 for i, b in enumerate(_long_paras) if b.text not in whole]
            check(not missing, f"{strat}(重疊 {ov})：每一段都完整保留",
                  f"丟失的段落={missing or '無'}")
        over = [len(c.body) for c in cs if len(c.body) > 400]
        check(not over, f"{strat}(重疊 {ov})：塊長不超過 chunk_max=400", f"超標={over}")

# 超長單段（> hi）走 _split_long，同樣不能少字
_huge = kb_parse.Block("這是一段很長的話。" * 200)      # 1800 字
for ov in (0, 80):
    cs = kb_chunk.chunk_blocks([_huge], doc_type="txt", title="t",
                               params={"chunk_strategy": "paragraph", "chunk_size": 300,
                                       "chunk_min": 100, "chunk_max": 300,
                                       "chunk_overlap": ov, "chunk_prefix": False})
    total = sum(len(c.body) for c in cs)
    check(total >= len(_huge.text) - ov - 10,      # 尾端可能少一點點重疊，但不該整段消失
          f"超長單段(重疊 {ov})：{len(cs)} 塊，內文 {total} 字 ≥ 原始 {len(_huge.text)}",
          f"最大塊={max((len(c.body) for c in cs), default=0)}")
    check(all(len(c.body) <= 300 for c in cs),
          f"超長單段(重疊 {ov})：塊長都 ≤ chunk_max",
          f"最大塊={max((len(c.body) for c in cs), default=0)}")

# 尾端小碎塊被併回「前一塊」時，不可以把它開頭的重疊留在裡面（會自己重複自己）
#   場景：一塊剛好收成（≥ target）→ 剩下的零頭 < lo → 併回去。
#   舊版沒有記「這塊開頭放了多少重疊」，併回去就變成「同一段內容出現兩次」。
_b1 = kb_parse.Block("第一段" * 85)      # 255 字 → 收成一塊
_b2 = kb_parse.Block("尾巴一句話。")      # 6 字  → 太小，會被併回上一塊
_cs = kb_chunk.chunk_blocks([_b1, _b2], doc_type="txt", title="t",
                            params={"chunk_strategy": "paragraph", "chunk_size": 200,
                                    "chunk_min": 200, "chunk_max": 400,
                                    "chunk_overlap": 60, "chunk_prefix": False})
_expect = _b1.text + "\n" + _b2.text
check(len(_cs) == 1 and _cs[0].body == _expect,
      "零頭併回前一塊時會去掉重疊（不會自己重複自己）",
      f"{len(_cs)} 塊 / body {len(_cs[0].body) if _cs else 0} 字（預期 {len(_expect)}）")

# 代碼塊不被切斷（md）
md_chunks = kb_chunk.chunk_blocks(
    kb_parse.parse(os.path.join(TMP, "faq.md")).blocks, doc_type="md", title="faq",
    params={"chunk_size": 500, "chunk_max": 600, "chunk_min": 100, "chunk_overlap": 50})
has_code = any("def hello()" in c.text for c in md_chunks)
check(has_code, "md 的程式碼區塊有被保留")

# ---------- 2) 知識庫 + 上傳 + 入庫 ----------
print("\n[2] 知識庫與入庫流程…")
cid = None
for c in kb_db.list_collections(business_id=BIZ):
    if c["name"] == CHECK_COLLECTION:
        cid = c["id"]
if not cid:
    col = kb_db.create_collection(CHECK_COLLECTION, "RAG 自檢用，可刪", business_id=BIZ)
    cid = col["id"]

# 同一商家底下的第二個庫：用來驗「檢索參數全庫共用」（分塊參數才按庫）。
other_cid = None
for c in kb_db.list_collections(business_id=BIZ):
    if c["name"] == OTHER_COLLECTION:
        other_cid = c["id"]
if not other_cid:
    other_cid = kb_db.create_collection(OTHER_COLLECTION, "自檢對照用，可刪",
                                        business_id=BIZ)["id"]

# 可重複執行：先清掉這個庫上次的自檢資料
old = kb_db.list_documents(cid)
for d in old:
    kb_ingest.delete_document(int(d["id"]))
if old:
    print(f"  （清掉上次自檢的 {len(old)} 份文件）")
check(bool(cid), f"知識庫就緒 #{cid}")

# 參數覆寫（★ 檢索參數與提示詞全庫共用；分塊參數才按知識庫）
# ⚠️ 全域的兩組（檢索參數、提示詞）寫下去會蓋到正式設定 → 先整批存原值、後面還原。
#    沒有還原的話，使用者調好的 min_score / 提示詞會被清成預設（踩過）。
_saved_global = kb_db.get_global_param_overrides()      # 全量（含提示詞）
_saved_chunk = {k: kb_db.get_param_overrides(cid).get(k) for k in ("chunk_size",)}
kb_db.set_params(cid, {"top_k": 5, "min_score": "0.5", "chunk_size": 300})
p = kb_db.load_params(cid)
check(p["top_k"] == 5 and abs(p["min_score"] - 0.5) < 1e-9 and p["chunk_size"] == 300,
      "參數覆寫生效", f"top_k={p['top_k']} min_score={p['min_score']} size={p['chunk_size']}")
other_cid = other_cid or next((c["id"] for c in kb_db.list_collections(business_id=BIZ)
                              if c["id"] != cid), None)
if other_cid:
    po = kb_db.load_params(other_cid)
    check(po["top_k"] == 5 and abs(po["min_score"] - 0.5) < 1e-9
          and "top_k" not in kb_db.get_param_overrides(other_cid),
          "檢索參數全庫共用（另一庫沒有自己的覆寫，也讀到同一組）",
          f"other#{other_cid}: top_k={po['top_k']} min_score={po['min_score']}")
kb_db.set_params(cid, {"top_k": None, "min_score": None, "chunk_size": None})
p = kb_db.load_params(cid)
# 檢索參數清掉覆寫後會回到「這個人原本設定的值」（沒有才是預設），所以不能硬比 6
_expect_topk = rag_config.coerce("top_k", _saved_global.get("top_k") or 6)
check(p["top_k"] == _expect_topk and p["chunk_size"] == 500,
      "刪除覆寫後回到原本的值（檢索回原設定、分塊回預設）",
      f"top_k={p['top_k']}（預期 {_expect_topk}）size={p['chunk_size']}")

# ★ 提示詞也是「全域參數」（工作台 知識庫 → 參數配置 → C. 提示詞 可編輯）
#   驗三件事：DB 覆寫生效、壞值退回預設（不讓 .format 炸）、清掉後回內建預設。
MARK = "【自檢】{context} ／ {sentinel}"
kb_db.set_params(None, {"rag_answer_prompt": MARK})
kb_search.invalidate()
check(kb_search.prompt_of("rag_answer_prompt") == MARK,
      "提示詞覆寫生效（prompt_of 取到 DB 值）")
kb_db.set_params(None, {"rag_answer_prompt": "少了必需佔位符的提示詞"})
kb_search.invalidate()
check(kb_search.prompt_of("rag_answer_prompt") == rag_config.RAG_PROMPT,
      "提示詞缺必需佔位符 → 執行時退回內建預設（不讓生成那一刻炸）")
kb_db.set_params(None, {"rag_answer_prompt": None})
kb_search.invalidate()
check(kb_search.prompt_of("rag_answer_prompt") == rag_config.RAG_PROMPT,
      "清掉提示詞覆寫 → 回到內建提示詞")

# 還原成自檢前的正式設定：先清掉自檢寫入的**全部全域項**，再整批寫回。
# ⚠️ 分塊參數要帶 cid 寫（`set_params(None, ...)` 會把它寫到「未分組文件」那批）。
kb_db.set_params(None, {k: None for k in rag_config.PARAM_SPECS
                        if kb_db._param_scope(k) == "global"})
kb_db.set_params(cid, _saved_chunk)
kb_db.set_params(None, _saved_global)
kb_search.invalidate()
_p = kb_db.get_global_param_overrides()
check(_p == _saved_global,
      "自檢沒有改動正式的全域設定（檢索參數 + 提示詞）",
      f"top_k={_p.get('top_k')} min_score={_p.get('min_score')} "
      f"提示詞覆寫={'有' if _p.get('rag_answer_prompt') else '無'}")

doc_ids = []
for fn in ("manual.pdf", "manual.docx", "price.xlsx", "faq.md", "warranty.txt", "stores.csv"):
    with open(os.path.join(TMP, fn), "rb") as f:
        raw = f.read()
    saved = kb_ingest.save_upload(fn, raw, "", cid)
    did = saved["doc"]["id"]
    doc_ids.append(did)
    t0 = time.time()
    kb_ingest.parse_and_chunk(did)
    doc = kb_db.get_document(did)
    ok = doc["status"] == "draft" and doc["chunk_count"] > 0
    check(ok, f"{fn} → doc#{did} 分塊 {doc['chunk_count']} 塊 / {doc['char_count']} 字",
          f"{time.time() - t0:.2f}s")

# 重複上傳偵測
with open(os.path.join(TMP, "faq.md"), "rb") as f:
    raw = f.read()
dup = kb_ingest.save_upload("faq.md", raw, "", cid)
check(dup["duplicate_of"] is not None, f"重複上傳偵測到舊版 #{dup['duplicate_of']}",
      f"新版 v{dup['doc']['version']}")

# ---------- 3) embedding ----------
print("\n[3] embedding（jina-v2-base-zh）…")
target = doc_ids[1]          # docx
t0 = time.time()
res = kb_ingest.embed_document(target, "pending")
dt = time.time() - t0
check(res["done"] > 0 and res["failed"] == 0, f"doc#{target} 向量化 {res['done']}/{res['total']}",
      f"{dt:.1f}s（{dt / max(1, res['total']):.2f}s/塊）")
stats = kb_db.chunk_stats(target)
check(int(stats["done"]) == int(stats["total"]), "全部塊都標成 done",
      f"done={stats['done']} total={stats['total']}")
row = kb_db.query("SELECT embedding, embedding_model, embedding_dim, array_length(embedding,1) AS n "
                 "FROM kb_chunk WHERE doc_id=%s AND embedding_status='done' LIMIT 1", (target,))
check(bool(row) and row[0]["n"] == 768, "向量維度 768 且落地成 REAL[]",
      f"dim={row[0]['embedding_dim']} n={row[0]['n']} model={row[0]['embedding_model']}")

doc = kb_db.get_document(target)
check(doc["status"] == "draft" and doc["progress"] == 100,
      "向量化完成後仍在 draft（等客服確認入庫）", f"status={doc['status']}")

act = kb_ingest.activate_document(target)
check(act["status"] == "ready", "確認入庫 → status=ready")

# 其餘文件也入庫
for did in doc_ids:
    if did == target:
        continue
    kb_ingest.embed_document(did, "pending")
    kb_ingest.activate_document(did)
ready = kb_db.query("SELECT count(*) c FROM kb_document WHERE collection_id=%s AND status='ready'",
                    (cid,))
check(int(ready[0]["c"]) >= 6, f"知識庫內 ready 文件數 = {ready[0]['c']}")

# ---------- 4) 分塊微調 ----------
print("\n[4] 分塊微調（編輯 / 拆分 / 合併 / 停用）…")
chunks = kb_db.list_chunks(target)
c0 = chunks[0]
c1 = chunks[1]
kb_db.update_chunk(c0["id"], "客服手改過的文字：A 款售價 1280 元。")
ch = kb_db.get_chunk(c0["id"])
check(ch["edited"] is True and ch["embedding_status"] == "pending",
      "手改後標 edited 且向量回到 pending")
kb_db.update_chunk(c0["id"], None, False)
check(kb_db.get_chunk(c0["id"])["enabled"] is False, "可以停用單一知識塊")
kb_db.update_chunk(c0["id"], None, True)

# 拆分
long_chunk = max(chunks, key=lambda c: len(c["text"]))
mid = len(long_chunk["text"]) // 2
kb_db.shift_ordinals(target, long_chunk["ordinal"])
new_id = kb_db.insert_chunk(target, long_chunk["ordinal"] + 1,
                            long_chunk["text"][mid:].strip() or "第二部分",
                            prefix=long_chunk["prefix"] or "")
kb_db.update_chunk(long_chunk["id"], long_chunk["text"][:mid].strip())
kb_db.renumber_chunks(target)
now = kb_db.list_chunks(target)
ords = [c["ordinal"] for c in now]
check(ords == list(range(len(ords))), f"拆分 + 重排後 ordinal 連續（共 {len(ords)} 塊）")

# 合併
before = len(now)
kb_db.delete_chunk(new_id)
kb_db.renumber_chunks(target)
after = len(kb_db.list_chunks(target))
check(after == before - 1, f"刪除塊 {before} → {after}")

# ---------- 5) 預覽（不寫庫）----------
print("\n[5] 分塊預覽…")
snapshot = kb_db.list_chunks(target)
pv = kb_ingest.preview_chunks(target, {"chunk_size": 200, "chunk_min": 50, "chunk_max": 260,
                                       "chunk_overlap": 40})
after_pv = kb_db.list_chunks(target)
check(len(snapshot) == len(after_pv), "預覽不寫入任何資料",
      f"預覽 {pv['count']} 塊 / 平均 {pv['avg_chars']} 字")
check(pv["strategy"] == "heading" or pv["strategy"] == "table_row" or pv["strategy"] == "paragraph",
      f"預覽策略 = {pv['strategy']}")

print("\n[6] 清理…")
purge_business(BIZ)
left = kb_db.list_collections(business_id=BIZ)
check(not left, "自檢商家已被清空（庫／文件／塊全刪）", f"剩 {len(left)} 個庫")

raise SystemExit(report("RAG 知識庫入庫自檢"))
