#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""知識庫「圖文」自檢：PDF 行級解析 + 圖片抽取/題注/綁定 + 圖片端點 + 命中帶圖。

涵蓋 2026-10-12 新增的三件事：
  ① PDF 改走**行級解析**（pdfplumber）：頁碼/頁眉頁腳會被丟掉、
     **跨頁只是物理隔斷**（不產生段落分隔）。
  ② 圖片抽取（`ParsedAsset`）＋題注優先＋裝飾圖過濾。
  ③ 圖**必須綁到知識塊**，命中時一起帶出來（圖文結合輸出）。

⚠️ 合成 PDF 用 **英文語料**：內建的 Type1 Helvetica 沒有中文字型，
   餵中文會被抽成 `(cid:xxx)` 亂碼（中文的真實案例請看 `kb_check.py` 的 docx）。

可重複執行（自建臨時知識庫，跑完刪乾淨）。用法：
    python tests/kb_asset_check.py
"""
import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 專案根
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from fastapi.testclient import TestClient

import main
from app.rag import kb_chunk
from app.rag import kb_db
from app.rag import kb_ingest
from app.rag import kb_parse
from app.rag import kb_search

TMP = os.path.join(HERE, "_p1_samples")
os.makedirs(TMP, exist_ok=True)
# ⚠️ 名字要跟「使用者自己在用的庫」分開 —— 這個庫會被建立／刪除。
CHECK_COLLECTION = "圖文自檢庫(測試)"
FAIL = []


def check(cond, label, extra=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {label}" + (f"   {extra}" if extra else ""))
    if not cond:
        FAIL.append(label)
    return cond


# ---------- 造檔 ----------

def make_pdf_pages(path, pages, w=595, h=842):
    """pages = [[(y, text), ...], ...]；y 是**從頁面底部算**的基線座標。

    ⇒ 要讓某行落在「頁底」（跨頁續接的判斷依據），y 要小（例如 95）；
      要落在「頁首」，y 要大（例如 780）。
    """
    def esc(s):
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    n = len(pages)
    objs, kids = {}, []
    for i, lines in enumerate(pages):
        pno, cno = 4 + 2 * i, 5 + 2 * i
        kids.append(f"{pno} 0 R")
        content = "BT /F1 12 Tf\n"
        for (y, text) in lines:
            content += f"1 0 0 1 50 {y} Tm ({esc(text)}) Tj\n"
        content += "ET"
        objs[pno] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {w} {h}] "
                     f"/Resources << /Font << /F1 3 0 R >> >> /Contents {cno} 0 R >>")
        objs[cno] = f"<< /Length {len(content)} >>\nstream\n{content}\nendstream"
    objs[1] = "<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {n} >>"
    objs[3] = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    out = b"%PDF-1.4\n"
    offs = {}
    for i in sorted(objs):
        offs[i] = len(out)
        out += f"{i} 0 obj\n{objs[i]}\nendobj\n".encode()
    total = max(objs) + 1
    x = len(out)
    out += f"xref\n0 {total}\n0000000000 65535 f \n".encode()
    for i in range(1, total):
        out += (f"{offs[i]:010d} 00000 n \n" if i in offs else "0000000000 65535 f \n").encode()
    out += f"trailer\n<< /Size {total} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)


def make_docx_with_image(path, big=(800, 500), small=(40, 30), caption="figure 1-1 setup page"):
    from PIL import Image
    import docx
    from docx.shared import Inches

    p1 = os.path.join(TMP, "_img_big.png")
    p2 = os.path.join(TMP, "_img_small.png")
    Image.new("RGB", big, (220, 235, 250)).save(p1)
    Image.new("RGB", small, (255, 255, 255)).save(p2)

    d = docx.Document()
    d.add_heading("操作說明", level=1)
    d.add_paragraph("第一步：開啟設定頁面。")
    d.add_picture(p1, width=Inches(4))
    d.add_paragraph(caption)                 # ← 題注
    d.add_paragraph("第二步：輸入連線資訊。")
    d.add_picture(p2)                        # ← 太小 → 應被過濾
    d.add_paragraph("第三步：儲存。")
    d.save(path)
    return path


# ======================================================================
# [1] PDF 行級解析
# ======================================================================

print("=" * 74)
print(" 知識庫圖文自檢（PDF 行級解析 + 圖片抽取/綁定）")
print("=" * 74)

print("\n[1] PDF 行級解析：頁碼／頁眉頁腳過濾 + 跨頁接續")
pdf_path = os.path.join(TMP, "linelevel.pdf")
make_pdf_pages(pdf_path, [
    # p1：標頭 + 一句話，**最後一行在頁底**（y 小）
    [(780, "Product manual introduction."),
     (760, "The A100 headset costs 1280 and ships in"),
     (95, "two business days from our warehouse")],
    # p2：**第一行在頁首**（y 大）→ 應該跟 p1 最後一行接起來
    [(780, "and is covered by a 12 month warranty."),
     (760, "Contact support for details."),
     (95, "- 2 -")],                                    # ← 頁碼（頁腳帶 + 頁碼樣式）
    [(780, "Spare parts can be ordered separately."),
     (95, "- 3 -")],
])
parsed = kb_parse.parse(pdf_path, "pdf")

check(parsed.meta.get("pages") == 3, "讀到 3 頁", str(parsed.meta.get("pages")))
check(parsed.meta.get("dropped_header_footer") == 2,
      "兩個頁碼都被辨識並丟掉",
      f"dropped={parsed.meta.get('dropped_header_footer')}")
check(not any(("- 2 -" in b.text) or ("- 3 -" in b.text) or b.text.strip().isdigit()
              for b in parsed.blocks),
      "頁碼沒有混進任何知識塊",
      str([b.text for b in parsed.blocks if b.text.strip().isdigit()]))
check(parsed.meta.get("cross_page_joins") == 1,
      "跨頁續接被辨識（p1 尾 → p2 首 只有 1 處）",
      f"joins={parsed.meta.get('cross_page_joins')}")
joined = [b.text for b in parsed.blocks if "warehouse and is covered" in b.text]
check(bool(joined), "跨頁那句被接成同一段（不是兩個段落）",
      str(joined)[:90])
check(joined and joined[0].endswith("warranty."),
      "接起來之後整句完整到句尾")
check(all(b.page_no is not None for b in parsed.blocks),
      "每個塊都有 page_no（圖片才綁得上）",
      str([b.page_no for b in parsed.blocks]))
check(any(b.text.endswith(".") for b in parsed.blocks),
      "真正以句點結束的句子仍然自成一段（沒被過度合併）")
# 頁碼被丟掉，但**內文一字不少**
_all = "".join(b.text for b in parsed.blocks)
for needle in ("Product manual introduction.", "The A100 headset costs 1280",
               "Contact support for details.", "Spare parts can be ordered separately."):
    if needle not in _all:
        check(False, f"內文沒有遺漏：{needle[:30]}")

print("\n[2] 圖片抽取 + 題注優先 + 裝飾圖過濾")
docx_path = make_docx_with_image(os.path.join(TMP, "with_image.docx"))
pd = kb_parse.parse(docx_path, "docx")
check(len(pd.assets) == 2, f"抽出 2 張圖", f"實際 {len(pd.assets)}")
big = next((a for a in pd.assets if a.width >= 100), None)
small = next((a for a in pd.assets if a.width < 100), None)
check(big is not None and big.caption == "figure 1-1 setup page",
      "題注被當成預設描述（不必再花 AI）",
      repr(big.caption if big else None))
check(big is not None and not big.filtered_reason, "正常的圖沒有被過濾掉")
check(small is not None and bool(small.filtered_reason),
      "太小的圖被標成裝飾（仍保留記錄）", repr(small.filtered_reason if small else None))
check(all(a.anchor_block is not None for a in pd.assets),
      "每張圖都有綁定錨點（anchor_block）",
      str([a.anchor_block for a in pd.assets]))

# ======================================================================
# [3] 入庫：落地 + 去重 + 綁定分塊
# ======================================================================

print("\n[3] 入庫：圖片落地 + 綁定知識塊")
cid = None
for c in kb_db.list_collections():
    if c["name"] == CHECK_COLLECTION:
        cid = c["id"]
if not cid:
    cid = kb_db.create_collection(CHECK_COLLECTION, "圖文自檢用，可刪")["id"]
else:
    kb_db.update_collection(cid, {"enabled": True})     # 自檢庫一定要啟用

with open(docx_path, "rb") as f:
    up = kb_ingest.save_upload("with_image.docx", f.read(), "", cid)
did = up["doc"]["id"]
kb_ingest.parse_and_chunk(did)

assets = kb_db.list_assets(doc_id=did)
check(len(assets) == 2, f"入庫後 kb_asset 有 2 筆", f"實際 {len(assets)}")
check(all(a["rel_path"] for a in assets), "每筆都有落地路徑")
on_disk = [a for a in assets
           if os.path.exists(os.path.join(HERE, "media", a["rel_path"]))
           or os.path.exists(os.path.join(os.getcwd(), "media", a["rel_path"]))]
check(len(on_disk) == len(assets), "圖檔真的寫到磁碟（media/ 底下）",
      f"{len(on_disk)}/{len(assets)}")
check(all(a["doc_id"] == did for a in assets), "圖掛在這份文件底下")
check(any(a["caption_source"] == "doc" for a in assets),
      "題注來源標成 doc")
check(any(not a["enabled"] and a["filtered_reason"] for a in assets),
      "被過濾的圖 enabled=false（不會被帶出去）")
check(all((a["chunk_count"] or 0) > 0 for a in assets if a["enabled"]),
      "**啟用中的圖都綁到了知識塊**（圖文結合輸出的前提）",
      str([(a["id"], a["chunk_count"]) for a in assets]))

# 重複入庫不該產生重複的圖（sha256 去重）
kb_ingest.parse_and_chunk(did)
check(len(kb_db.list_assets(doc_id=did)) == 2,
      "重新分塊不會產生重複圖片（sha256 去重）",
      f"{len(kb_db.list_assets(doc_id=did))} 筆")

# ======================================================================
# [4] 圖片端點
# ======================================================================

print("\n[4] 圖片端點（列表 / 人工描述 / AI 不覆蓋人工 / 停用不帶出）")
client = TestClient(main.app)


def body(res):
    import json
    raw = getattr(res, "content", None) or getattr(res, "body", b"")
    return json.loads(raw.decode("utf-8"))


r = client.get(f"/console/kb/assets?doc_id={did}&include_filtered=1")
d = body(r)
check(r.status_code == 200 and len(d.get("assets", [])) == 2,
      "GET /assets 回 2 張", f"status={r.status_code}")
check(d.get("stats", {}).get("with_caption", 0) >= 1, "統計：有描述的張數 > 0",
      str(d.get("stats")))
r2 = client.get(f"/console/kb/assets?doc_id={did}&include_filtered=0")
check(len(body(r2).get("assets", [])) == 1, "include_filtered=0 只回內容圖")

aid = next(a["id"] for a in assets if a["caption_source"] == "doc")
rf = client.get(f"/console/kb/asset/file/{aid}")
check(rf.status_code == 200 and len(rf.content) > 100,
      "GET /asset/file/{id} 回得到圖檔位元組", f"{len(rf.content)} bytes")
check(client.get("/console/kb/asset/file/999999").status_code == 404,
      "不存在的圖回 404")

r = client.post("/console/kb/asset/caption", json={"id": aid, "caption": "人工寫的描述"})
check(r.status_code == 200 and body(r)["asset"]["caption_source"] == "manual",
      "人工描述寫入，來源標 manual")

# ⚠️ 這條**不會真的呼叫模型**：端點在 manual 時直接跳過 → 可以放心跑
r = client.post("/console/kb/asset/caption/ai", json={"id": aid})
rr = (body(r).get("results") or [{}])[0]
check(rr.get("skipped") and not rr.get("ok"),
      "AI 描述**不會覆蓋**人工寫的描述（要覆蓋得帶 force=1）",
      str(rr.get("error"))[:50])
check(kb_db.get_asset(aid)["caption"] == "人工寫的描述",
      "人工描述沒有被動到")

r = client.post("/console/kb/asset/caption", json={"id": aid, "caption": ""})
check(body(r)["asset"]["caption_status"] == "pending", "清空描述 → 回到待填")

# 停用後不該被檢索帶出
small_id = next(a["id"] for a in assets if not a["enabled"])
check(small_id not in [x["id"] for x in
                       (kb_db.assets_for_chunks([c["id"] for c in kb_db.list_chunks(did)]).get(
                           kb_db.list_chunks(did)[0]["id"]) or [])],
      "被停用的圖不會出現在 assets_for_chunks")

# ======================================================================
# [5] 命中帶圖
# ======================================================================

print("\n[5] 檢索命中 → 連圖一起帶出來")
res_embed = kb_ingest.embed_document(did, "pending")
kb_ingest.activate_document(did)
kb_search.invalidate()
check(res_embed.get("done", 0) > 0, "向量化完成", str(res_embed))

sr = kb_search.search("設定頁面怎麼開啟")
check(sr.passed, "檢索命中", f"top1={sr.top1:.3f}")
check(bool(sr.assets), "SearchResult 帶出了圖片", f"{len(sr.assets)} 張")
check(all("url" in a for a in sr.assets), "每張圖都有取圖 url")
hit_with_assets = [h for h in sr.hits if h.assets]
check(bool(hit_with_assets), "至少有一個入選塊帶著圖（Hit.assets）")
js = sr.hits[0].as_dict() if sr.hits else {}
check("assets" in js, "Hit.as_dict() 含 assets（寫進 matched_rows，紀錄頁才看得到）")

# ======================================================================
# [6] 圖片內嵌 markdown（2026-10-12）：圖片 → chunk 文字 / 只對文字向量化
# ======================================================================

print("\n[6] 圖片內嵌 markdown + 圖片庫可被檢索")
from app.rag import kb_chunk as _kc          # noqa: E402
from app.rag import rag_config as _rc        # noqa: E402

kb_ingest.embed_document(did, "pending")
kb_ingest.activate_document(did)
kb_search.invalidate()

chunks = kb_db.list_chunks(did)
withmd = [c for c in chunks if "![" in (c["text"] or "")]
check(bool(withmd), "chunk 文字裡有內嵌的 markdown 圖片標記",
      (withmd[0]["text"][-90:] if withmd else "（沒有）"))
check(any("/media/" in (c["text"] or "") for c in chunks),
      "標記用的是公開路徑 /media/…（前端直接渲染）")
check(any("figure 1-1 setup page" in (c["text"] or "") for c in chunks),
      "標記的 alt 用了題注文字（沒題注才留空）")

check(_kc.strip_md_images("前\n![圖 1](/media/a.png)\n後") == "前\n\n後",
      "strip_md_images 剝掉圖片標記只留文字（向量化只吃文字）")
check(_kc.strip_md_images("沒有圖片的文字") == "沒有圖片的文字", "沒有標記時原樣回傳")

a_on = [a for a in kb_db.list_assets(doc_id=did) if a["enabled"]]
check(bool(a_on) and all(_rc.media_url_of_rel(a["rel_path"]).startswith("/media/")
                         for a in a_on),
      "啟用中的圖都有公開網址 /media/…",
      [_rc.media_url_of_rel(a["rel_path"]) for a in a_on])
check(all(os.path.exists(os.path.join(_rc.RAG_MEDIA_DIR, os.path.basename(a["rel_path"])))
          for a in a_on),
      "圖檔真的落在公開目錄 media/kb_assets/（不是整包 media 都公開）")

imgcol = kb_db.get_image_collection(cid)
check(bool(imgcol) and imgcol["kind"] == "image" and imgcol["image_of"] == cid,
      "文檔庫有對應的圖片庫（kind=image / image_of=cid）",
      (f"#{imgcol['id']} {imgcol['name']}" if imgcol else "（沒有）"))
check(bool(imgcol) and all(a["collection_id"] == imgcol["id"] for a in a_on),
      "從文件抽出的圖落在該文檔庫的圖片庫",
      [(a["id"], a["collection_id"]) for a in a_on])

# 寫描述（模擬真實流程：描述寫好 → 立刻向量化 → 之後就查得到）
own_aid = a_on[0]["id"] if a_on else None
r = client.post("/console/kb/asset/caption",
                json={"id": own_aid, "caption": "figure 1-1 setup page"})
check(r.status_code == 200 and body(r)["asset"]["caption_source"] == "manual",
      "圖片描述寫入成功（來源 manual）")
check(kb_db.get_asset(own_aid)["embedding_status"] == "done",
      "★ 寫完描述**立刻向量化**（圖片能被檢索的前提；以前完全沒這一步）")
kb_search.invalidate()

r = kb_search.search("figure 1-1 setup page")
imhits = [h for h in r.candidates if h.kind == "image"]
check(bool(imhits), "★ 圖片庫條目進得了候選池（描述可被檢索）— 以前完全搜不到",
      [(h.kind, round(h.final, 3)) for h in r.candidates[:3]])
own = [h for h in imhits if h.asset_id == own_aid]
check(bool(own), "這張圖的條目在候選池裡（text 就是 markdown，前端渲染即可）",
      (own[0].text if own else [h.text[:44] for h in imhits[:3]]))
check(bool(own) and own[0].text.startswith("![figure 1-1"),
      "圖片命中的 text 的 alt 用了描述", (own[0].text[:60] if own else ""))

# 去重：綁在入選塊上的圖，不該同時以「圖片條目」的身分再出現一次
bound_ids = set()
for _arr in (kb_db.assets_for_chunks([h.chunk_id for h in r.hits
                                      if h.kind != "image"]).values() or []):
    bound_ids |= {int(x["id"]) for x in _arr}
dup = [h for h in r.hits if h.kind == "image" and h.asset_id in bound_ids]
check(not (bound_ids and dup),
      "同一張圖同時被 chunk 與圖片庫命中時只出一次",
      f"綁定中的圖 {sorted(bound_ids)} / 重複出現的圖片命中 {[h.asset_id for h in dup]}")


kb_db.set_collection_enabled(cid, False)
kb_search.invalidate()
r2 = kb_search.search("figure 1-1 setup page")
check(not [h for h in r2.candidates if h.kind == "image" and h.asset_id == own_aid],
      "文檔庫停用後它的圖片庫也檢索不到（啟停用連動）")
check(kb_db.get_collection(imgcol["id"])["enabled"] is False, "圖片庫 enabled 被連動關閉")
kb_db.set_collection_enabled(cid, True)
kb_search.invalidate()
check(bool([h for h in kb_search.search("figure 1-1 setup page").candidates
            if h.kind == "image" and h.asset_id == own_aid]), "恢復啟用後又能檢索到")

# ---------- [7] 送出端：文字與圖片只發一條（客戶不會收到一對重複的） ----------
#
# 症狀（使用者 實測）：知識庫命中帶圖時，客戶端收到
#   「一條不帶圖的純文字」+「一條一模一樣、但帶圖的」。
# 根因：舊版固定「文字一條 → 圖一張（caption 又是同一段文字）」。
# 修法：能帶 caption 的渠道（WhatsApp）把文字塞進圖片一起發，只發一條；
#       不能帶的（微信）維持文字 + 圖，但圖不帶文字、順序是「先說明後圖」。
print("\n[7] 送出端：圖文只發一條（不再「一條純文字 + 一條一模一樣帶圖的」）")

import asyncio

from app.autoreply import autoreply
from app.channels import igfb
from app.channels import whatsapp_send, wechat_send

_a = kb_db.get_asset(own_aid)
_asset = dict(_a or {})
_text = "伺服器端轉發成功的日誌如下圖所示。"

# 策略判斷（純函式）
check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset], _text) is True,
      "WhatsApp 私訊 + 有圖 + 有文字 → 走「圖文一條」")
check(autoreply.assets_inline_caption("wechat", "dm", [_asset], _text) is False,
      "微信圖片沒有 caption → 不宣稱能一條（平台限制，只能文字+圖）")
check(autoreply.assets_inline_caption("whatsapp", "comment", [_asset], _text) is False,
      "貼文留言不發圖（公開場合）")
check(autoreply.assets_inline_caption("whatsapp", "dm", [], _text) is False,
      "沒有圖 → 沒什麼好 inline 的")
check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset], "   ") is False,
      "只有空白文字 → 沒有東西可以當 caption")
check(autoreply.assets_inline_caption("whatsapp", "dm", [_asset], "字" * 2000) is False,
      "文字超過 caption 上限（1024）→ 改走「文字一條 + 圖一條」不硬塞")

# 實測 _send_assets：caption 只貼第一張、captioned 只在真的送出去時為 True
_shots = []          # [(caption, ...)] 每次 send_image 的 caption

async def _fake_upload(path, mime, fname):
    return "MEDIA_ID"

async def _fake_send_image(to, mid, caption=""):
    _shots.append(caption)
    return {"status": 200, "body": "{}"}

def _fake_send_image_wechat_sync(openid, local_path):
    _shots.append("（微信圖片不帶文字）")
    return {"status": 200, "body": "{}"}

_orig_up, _orig_img = whatsapp_send.upload_media, whatsapp_send.send_image
_orig_wx = wechat_send.send_image
# 兩個渠道都要換成假的：自檢**不准真的打外部 API**
whatsapp_send.upload_media, whatsapp_send.send_image = _fake_upload, _fake_send_image
wechat_send.send_image = _fake_send_image_wechat_sync
try:
    _shots.clear()
    st = asyncio.run(autoreply._send_assets("whatsapp", "dm", [_asset],
                                            to="85200000000", caption=_text))
    check(st["sent"] == 1 and st.get("captioned") is True,
          "WhatsApp 一張圖 → 送出 1、captioned=True",
          f"sent={st['sent']} captioned={st.get('captioned')}")
    check(_shots == [_text], "文字確實當成 caption 送出去了（只送一次）", repr(_shots[:2]))

    two = [dict(_asset), dict(_asset)]
    _shots.clear()
    st2 = asyncio.run(autoreply._send_assets("whatsapp", "dm", two,
                                             to="85200000000", caption=_text))
    check(st2["sent"] == 2 and st2.get("captioned") is True,
          "兩張圖 → 兩張都送出、captioned=True")
    check(_shots[0] == _text and _shots[1] == "",
          "★ caption 只貼第一張（否則客戶會看到同一段話重複出現）", repr(_shots))

    _shots.clear()
    st3 = asyncio.run(autoreply._send_assets("whatsapp", "dm", [_asset], to="85200000000"))
    check(st3.get("captioned") is False and _shots == [""],
          "沒給 caption → captioned=False（呼叫端才知道要自己補發文字）")

    _shots.clear()
    st4 = asyncio.run(autoreply._send_assets("wechat", "dm", [_asset], to="openid"))
    check(st4["sent"] == 1 and st4.get("captioned") is False,
          "★ 微信 → 圖送出但 captioned=False（圖片不帶文字），呼叫端才知道要補發文字",
          f"sent={st4['sent']}")

    _shots.clear()
    st5 = asyncio.run(autoreply._send_assets("whatsapp", "comment", [_asset],
                                             to="c1", caption=_text))
    check(st5["sent"] == 0 and st5.get("skipped") == 1, "公開留言不發圖（略過）")

    st6 = asyncio.run(autoreply._send_assets("whatsapp", "dm",
                                             [{"rel_path": "kb_assets/沒有這個檔.png"}],
                                             to="85200000000", caption=_text))
    check(st6["failed"] == 1 and st6.get("captioned") is False,
          "圖檔不存在 → 記 failed、captioned=False（呼叫端會補發文字）")

    # ---- IG / FB：Meta 會「自己來下載」那個網址，所以一定要公開絕對 URL ----
    # 症狀（使用者）：IG 客戶只收到文字、沒有圖。根因是這裡原本直接 skipped
    # （註解寫「本服務沒有公開圖床」—— 但其實有：/media + PUBLIC_BASE_URL）。
    calls = []

    async def _fake_send_dm(platform, recipient_id, message_type, text, file_url, mime):
        calls.append({"platform": platform, "type": message_type, "url": file_url,
                      "mime": mime, "text": text})
        return {"status": 200, "body": '{"message_id":"mid.test"}'}

    _orig_dm = igfb.send_dm
    igfb.send_dm = _fake_send_dm
    try:
        url = autoreply.public_asset_url(_asset["rel_path"])
        check(url.startswith("https://") and "/media/" in url,
              "public_asset_url() 拼出外網可達的絕對網址", url)
        check(autoreply.public_asset_url("kb/舊圖/assets/x.png") == "",
              "舊圖不在公開目錄 → 回空字串（呼叫端才知道發不出去）")

        calls.clear()
        st_ig = asyncio.run(autoreply._send_assets("ig", "dm", [_asset],
                                                   to="1000000000000000", caption=_text))
        check(st_ig["sent"] == 1, "★ IG 會真的把圖送出去（不再靜默 skipped）",
              f"sent={st_ig['sent']} reason={st_ig.get('reason')!r}")
        check(calls and calls[0]["type"] == "file" and calls[0]["url"].startswith("https://"),
              "IG 走附件型別、且用公開絕對 URL", str(calls[:1])[:110])
        check(st_ig["captioned"] is False,
              "IG 的 attachment 沒有 caption → captioned=False（呼叫端會補發文字）")

        calls.clear()
        st_fb = asyncio.run(autoreply._send_assets("fb", "dm", [_asset], to="psid.test"))
        check(st_fb["sent"] == 1 and calls[0]["platform"] == "fb", "FB 同樣走得通")

        # 沒設 PUBLIC_BASE_URL → 要優雅降級 + 講清楚原因（不是靜默）
        from app.core import config as _c
        _saved_base = _c.PUBLIC_BASE_URL
        _c.PUBLIC_BASE_URL = ""
        try:
            check(autoreply.public_asset_url(_asset["rel_path"]) == "",
                  "沒設 PUBLIC_BASE_URL → 拼不出絕對網址")
            st_nb = asyncio.run(autoreply._send_assets("ig", "dm", [_asset], to="x"))
            check(st_nb["skipped"] == 1 and "PUBLIC_BASE_URL" in (st_nb.get("reason") or ""),
                  "沒設定時 skipped 且 reason 說明原因", (st_nb.get("reason") or "")[:46])
        finally:
            _c.PUBLIC_BASE_URL = _saved_base
    finally:
        igfb.send_dm = _orig_dm

    check(autoreply.assets_inline_caption("ig", "dm", [_asset], _text) is False,
          "IG 不會被當成「文字塞進 caption 的一條」（它的 attachment 沒有 caption）")
finally:
    whatsapp_send.upload_media, whatsapp_send.send_image = _orig_up, _orig_img
    wechat_send.send_image = _orig_wx

# ---------- [8] 端到端：handle_inbound 一次命中只送一條 ----------
#
# 這是使用者 實際踩到的場景。把「引擎」與「送出」都換成假的（**不真的發任何訊息**），
# 只看 handle_inbound 到底叫了幾次送出 —— 舊版固定是「文字 1 次 + 圖片 1 次」。
print("\n[8] 端到端：handle_inbound 命中帶圖時只送一條（不再重複）")

from app.llm import reply_engine

_ANSWER_TEXT = "伺服器端轉發成功的日誌如下圖所示。\n![圖 1 截圖](/media/placeholder.png)"
_calls = []


def _install_fakes(with_wechat=False):
    """把 handle_inbound 的外部依賴全部換成假的；回傳 (還原函式, 記錄清單)。"""
    from app.autoreply import autoreply_db

    saved = {}
    for name in ("is_enabled", "_guard", "_send", "_send_assets", "_mark", "_log"):
        saved[("autoreply", name)] = getattr(autoreply, name)
    saved[("reply_engine", "answer")] = reply_engine.answer
    saved[("autoreply_db", "claim_message")] = autoreply_db.claim_message
    saved[("autoreply_db", "mark_forward_log_handled")] = autoreply_db.mark_forward_log_handled
    saved[("db", "insert_forward_log")] = autoreply.db.insert_forward_log

    async def _fake_answer(question, **kw):
        return reply_engine.Answer(reply_engine.TIER_RAG, text=_ANSWER_TEXT,
                                   assets=[dict(_asset)], detail="測試")

    async def _fake_send(ch, sc, text, *, to="", comment_id=""):
        _calls.append(("text", text))
        return {"ok": True, "status": 200, "body": "{}"}

    async def _fake_send_assets(ch, sc, assets, *, to="", caption=""):
        _calls.append(("image", caption))
        return {"sent": len(assets), "failed": 0, "skipped": 0,
                "captioned": bool(caption.strip()) and ch == "whatsapp"}

    async def _fake_insert(_payload):
        return None

    autoreply.is_enabled = lambda ch, sc: True
    autoreply._guard = lambda k, ignore=False: ""
    autoreply._send = _fake_send
    autoreply._send_assets = _fake_send_assets
    autoreply._mark = lambda *a, **kw: None
    autoreply._log = lambda payload, claim_key=None: None
    reply_engine.answer = _fake_answer
    autoreply_db.claim_message = lambda key, payload: True
    autoreply_db.mark_forward_log_handled = lambda ch, cid: None
    autoreply.db.insert_forward_log = _fake_insert

    def _restore():
        autoreply.is_enabled = saved[("autoreply", "is_enabled")]
        autoreply._guard = saved[("autoreply", "_guard")]
        autoreply._send = saved[("autoreply", "_send")]
        autoreply._send_assets = saved[("autoreply", "_send_assets")]
        autoreply._mark = saved[("autoreply", "_mark")]
        autoreply._log = saved[("autoreply", "_log")]
        reply_engine.answer = saved[("reply_engine", "answer")]
        autoreply_db.claim_message = saved[("autoreply_db", "claim_message")]
        autoreply_db.mark_forward_log_handled = saved[("autoreply_db", "mark_forward_log_handled")]
        autoreply.db.insert_forward_log = saved[("db", "insert_forward_log")]

    return _restore


def _run_inbound(channel):
    return asyncio.run(autoreply.handle_inbound(
        channel=channel, scope="dm", text="服务器端转发成功日志长啥样？",
        to="85200000000", conversation_key="__kb_asset_check__",
        ignore_cooldown=True))


_restore = _install_fakes()
try:
    _calls.clear()
    res = _run_inbound("whatsapp")
    kinds = [c[0] for c in _calls]
    check(res.get("state") == "hit", "handle_inbound 走命中分支", str(res)[:80])
    check(kinds == ["image"],
          "★ WhatsApp：只送出 1 條（圖片 + caption），不再多送一條純文字", repr(kinds))
    check(_calls and _calls[0][1].startswith("伺服器端轉發成功的日誌"),
          "那條的文字就是回覆內容（當成 caption）", repr(_calls[0][1])[:40] if _calls else "")

    _calls.clear()
    _run_inbound("wechat")
    kinds = [c[0] for c in _calls]
    check(kinds == ["text", "image"],
          "微信：文字一條 + 圖片一條（圖片沒有 caption，平台限制）", repr(kinds))
    check(_calls[1][1] == "",
          "★ 微信那張圖不帶文字 → 客戶不會看到重複的一段話")

    # IG：使用者 回報「ins 端只有文字」—— 舊版附圖那一步靜默 skip。
    # 現在應該看到「文字一條 + 圖一條」（IG 的 attachment 沒有 caption）。
    _calls.clear()
    _run_inbound("ig")
    kinds = [c[0] for c in _calls]
    check(kinds == ["text", "image"],
          "★ IG：文字 + 圖都會送（不再只有文字）", repr(kinds))
    check(_calls[1][1] == "", "IG 那張圖不帶文字（attachment 沒有 caption）")
finally:
    _restore()

# ---------- [9] 標題識別 + 表格結構化（2026-10-12，「分塊效果不好」的根因） ----------
#
# 症狀（使用者 回報 doc#319）：格式清晰的 PDF，分塊卻一團糟 ——
#   · 4 個圖片題注被當成章節標題，把後面 51/74 個 block 的 heading_path 全汙染了
#   · 12 個真正的編號標題（「1.1 产品简介」）一個都沒認出來
#   · 標題與正文黏成一行（「1.1 产品简介本智能办公助手系统是面向…」）
#   · 表格被壓成流水文字（「故障现象 报错代码 核心原因 解决方案」後面接著糊掉的內容）
print("\n[9] 標題識別 + 表格結構化（分塊品質的關鍵）")

# 編號式章節標題：要認得出來，也要估對層級
for t, lv in [("1. 产品概述", 1), ("3.2 完整安装步骤（图文对应）", 2),
              ("1.1.2 细节说明", 3), ("5. 常见故障与解决方案（RAG 高频问答库）", 1)]:
    got = kb_parse._numbered_heading_level(t)
    check(got == lv, f"編號標題「{t}」→ 第 {lv} 層", f"實際 {got}")

# ★ 這些**不能**當標題：列表項（含句中標點）、版本號（數字後沒空白）
for t in ["1. 关闭电脑防火墙及杀毒软件拦截，避免安装包解压、组件注册被拦截",
          "2. 首次登录强制修改密码，密码需包含大小写字母、数字、8位及以上",
          "11.0及以上",
          "4. 初始化完成后自动进入主功能界面，可正常使用检索、问答功能"]:
    got = kb_parse._numbered_heading_level(t)
    check(got == 0, f"不算標題「{t[:20]}…」", f"實際 {got}")


def _ln(text, **kw):
    d = {"text": text, "x0": 90, "x1": 300, "top": 200, "bottom": 212}
    d.update(kw)
    return d


check(kb_parse._looks_heading(_ln("圖 2 文件解压完成目录示意图"), 595) is False,
      "★ 圖片題注不當標題（它由 _match_caption 綁到圖上）")
check(kb_parse._looks_heading(_ln("1.1 产品简介"), 595) is True,
      "★ 左對齊的編號標題要認得出來（原本只認置中 → 全漏）")
check(kb_parse._looks_heading(_ln("11.0及以上", x0=228, x1=284, in_table=True), 595) is False,
      "★ 表格內的欄位文字不當標題（偏右會像置中）")
check(kb_parse._looks_heading(_ln("第一章 总则"), 595) is True, "「第X章」仍認得")
check(kb_parse._heading_level_of("第三章 总则") == 1, "_heading_level_of 認得第X章")

# 表格 → markdown（列／欄關係要保住）
md = kb_parse._rows_to_markdown([["故障现象", "报错代码"], ["软件安装完成后无\n法启动", "E001"]])
check("| 故障现象 | 报错代码 |" in md, "表格轉成 markdown 表頭", repr(md.split("\n")[0]))
check("| 软件安装完成后无法启动 | E001 |" in md,
      "★ 儲存格內換行接回同一句（中文不加空格，否則「无法启动」檢索不到）",
      repr(md.split("\n")[-1]))
check(kb_parse._rows_to_markdown([["只有一欄"], ["x"]]) == "",
      "只有一欄的「表」判為誤判 → 不轉", "（免得把版心框當表格）")
check(kb_parse._rows_to_markdown([]) == "" and kb_parse._rows_to_markdown(None) == "",
      "空表格不炸")

# 前綴／內文去重（同一句別出現兩次）
check(kb_chunk._same_as_head("RAG问答测试专用｜图文版用户说明文档",
                             "RAG 问答测试专用｜图文版用户说明文档\n文档用途：…"),
      "前綴與內文開頭重複（忽略空白）→ 判為重複")
check(not kb_chunk._same_as_head("產品手冊", "耳機的保固期是 12 個月"),
      "不相干的前綴與內文 → 不判重複")
check(kb_chunk._drop_dup_heading("3.3 首次登录初始化设置\n软件首次打开需完成账号登录",
                                 "3. 软件安装 › 3.3 首次登录初始化设置")
      == "软件首次打开需完成账号登录",
      "★ 內文開頭的章節標題會被去掉（標題已在 prefix 裡）")
check(kb_chunk._drop_dup_heading("3.3 首次登录初始化设置", "3.3 首次登录初始化设置")
      == "3.3 首次登录初始化设置",
      "整塊就只有那個標題時 → 保留（去掉就空了）")

# 收尾要把圖片庫也一起清掉（它不屬於 kb_document 的 cascade）
_ick = (imgcol or {}).get("id")

# ---------- 收尾 ----------
print()
kb_ingest.delete_document(did)
if _ick:
    for _a in kb_db.list_assets(collection_id=_ick):
        kb_db.delete_asset(_a["id"])
    kb_db.delete_collection(_ick)
left = kb_db.asset_stats(doc_id=did).get("total") or 0
check(left == 0, "刪文件時圖片一起清掉", f"殘留 {left}")

# 自檢庫用完就刪（不要在使用者的清單裡留垃圾）
kb_db.delete_collection(cid)

print("\n" + "=" * 74)
if FAIL:
    print(f" 結果：{len(FAIL)} 項失敗")
    for f in FAIL:
        print("   ✗", f)
else:
    print(" 結果：全部通過 ✅")
print("=" * 74)
sys.exit(1 if FAIL else 0)
