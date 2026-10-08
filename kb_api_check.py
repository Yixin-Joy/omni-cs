#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""RAG 知識庫（P1 端點）自檢：用 TestClient 真打 /console/kb/*。

可重複執行：會自建臨時知識庫、跑完自己刪乾淨。
（入庫管道本身請用 `kb_check.py`。）

用法：
    python kb_api_check.py
"""
import json
import os
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from fastapi.testclient import TestClient
import main
import kb_db

FAIL = []
TMP = os.path.join(HERE, "_p1_samples")


def check(cond, label, extra=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {label}" + (f"   {extra}" if extra else ""))
    if not cond:
        FAIL.append(label)
    return cond


def body(res):
    """TestClient 回的是 httpx.Response（用 .content），不是 Starlette Response（.body）。"""
    raw = getattr(res, "content", None)
    if raw is None:
        raw = getattr(res, "body", b"")
    try:
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return {}


def ensure_samples() -> list[str]:
    """自建最小測試檔（不依賴 kb_check.py，兩支腳本各自可獨立執行）。"""
    os.makedirs(TMP, exist_ok=True)
    names = []

    p = os.path.join(TMP, "api_manual.md")
    with open(p, "w", encoding="utf-8") as f:
        f.write("# 商品說明\n\n## 藍牙耳機\nA 款售價 1280 元，續航 30 小時。\n"
                "B 款售價 890 元，支援主動降噪。\n\n## 退貨\n到貨 7 天內包裝完整可無理由退貨。\n")
    names.append("api_manual.md")

    p = os.path.join(TMP, "api_price.csv")
    with open(p, "w", encoding="utf-8") as f:
        f.write("產品,價格,庫存\n")
        for i in range(1, 21):
            f.write(f"耳機-{i},{800 + i * 13},{50 + i}\n")
    names.append("api_price.csv")

    try:
        import docx
        p = os.path.join(TMP, "api_doc.docx")
        d = docx.Document()
        d.add_heading("保固條款", level=1)
        d.add_paragraph("電子產品自購買日起提供一年原廠保固，人為損壞需自費維修。")
        d.save(p)
        names.append("api_doc.docx")
    except Exception as e:
        print(f"  （略過 docx 樣本：{e}）")
    return names


with TestClient(main.app) as c:
    print("=" * 72)
    print(" RAG 知識庫端點自檢（TestClient 真打 /console/kb/*）")
    print("=" * 72)

    print("\n[1] 概覽 / 集合")
    r = c.get("/console/kb/overview")
    d = body(r)
    check(r.status_code == 200 and "collections" in d, "GET /overview",
          f"model={d.get('model', {}).get('model')} downloaded={d.get('model', {}).get('downloaded')}")
    check(bool(d.get("model", {}).get("max_tokens")), "模型資訊含 token 上限",
          f"max_tokens={d.get('model', {}).get('max_tokens')}")

    r = c.post("/console/kb/collections/create", json={"name": "端點驗證庫", "description": "t"})
    d = body(r)
    cid = (d.get("collection") or {}).get("id")
    check(r.status_code == 200 and cid, f"POST /collections/create → #{cid}")

    r = c.post("/console/kb/collections/create", json={"name": "  "})
    check(r.status_code == 400, "空名稱被擋下", body(r).get("error"))

    r = c.post("/console/kb/collections/update", json={"id": cid, "description": "更新後"})
    check(r.status_code == 200 and body(r)["collection"]["description"] == "更新後",
          "POST /collections/update")

    print("\n[2] 參數配置")
    r = c.get(f"/console/kb/params?collection_id={cid}")
    d = body(r)
    keys = [i["key"] for i in d.get("items", [])]
    check(r.status_code == 200 and "top_k" in keys and "chunk_size" in keys,
          f"GET /params 回 {len(keys)} 個參數")
    check(len(d.get("retrieval", [])) == 6 and len(d.get("chunk", [])) == 6,
          "參數分組正確（retrieval 6 / chunk 6）")

    r = c.post("/console/kb/params", json={"collection_id": cid,
                                          "values": {"top_k": 99, "min_score": "-5",
                                                     "chunk_strategy": "不存在"}})
    d = body(r)
    vals = {i["key"]: i["value"] for i in d["params"]["items"]}
    check(vals["top_k"] == 30 and vals["min_score"] == 0.0 and vals["chunk_strategy"] == "auto",
          "超出範圍被夾到邊界 / 非法選項回預設",
          f"top_k={vals['top_k']} min_score={vals['min_score']} strategy={vals['chunk_strategy']}")
    c.post("/console/kb/params", json={"collection_id": cid,
                                       "values": {"top_k": None, "min_score": None,
                                                  "chunk_strategy": None}})

    print("\n[3] 上傳 → 解析 → 分塊")
    samples = ensure_samples()
    files = []
    for fn in samples:
        with open(os.path.join(TMP, fn), "rb") as f:
            files.append(("files", (fn, f.read(), "application/octet-stream")))
    r = c.post("/console/kb/upload", files=files, data={"collection_id": str(cid)})
    d = body(r)
    ok_all = all(x.get("ok") for x in d.get("results", []))
    doc_ids = [x["doc_id"] for x in d.get("results", []) if x.get("ok")]
    check(r.status_code == 200 and ok_all and len(doc_ids) == len(samples),
          f"POST /upload {len(samples)} 檔 → doc {doc_ids}")

    # 等背景 parse+chunk 完成
    for _ in range(40):
        docs = body(c.get(f"/console/kb/documents?collection_id={cid}"))["documents"]
        if all(x["chunk_count"] > 0 or x["status"] == "failed" for x in docs):
            break
        time.sleep(0.25)
    docs = body(c.get(f"/console/kb/documents?collection_id={cid}"))["documents"]
    check(all(x["chunk_count"] > 0 for x in docs),
          "背景分塊完成", str([(x["filename"], x["status"], x["chunk_count"]) for x in docs]))

    # 不支援格式
    r = c.post("/console/kb/upload",
               files=[("files", ("bad.xyz", b"hello", "application/octet-stream"))],
               data={"collection_id": str(cid)})
    e = body(r)["results"][0]
    check(not e["ok"] and "格式" in e["error"], "不支援格式給人話原因", e["error"])

    print("\n[4] 文件詳情 / 預覽 / 重新分塊")
    did = doc_ids[0]
    r = c.get(f"/console/kb/document?id={did}&chunks=1")
    d = body(r)
    check(r.status_code == 200 and d["doc"]["id"] == did and len(d["chunks"]) > 0,
          f"GET /document?id={did} 回 {len(d['chunks'])} 塊")

    r = c.post("/console/kb/document/preview",
               json={"id": did, "params": {"chunk_size": 120, "chunk_min": 40, "chunk_max": 150}})
    d = body(r)
    check(r.status_code == 200 and d["count"] > 0, "POST /document/preview",
          f"{d['count']} 塊 / 平均 {d['avg_chars']} 字 / 策略 {d['strategy']}")

    r = c.post("/console/kb/document/reparse", json={"id": did})
    check(r.status_code == 200, "POST /document/reparse（背景）", str(body(r)))
    # 等背景重新分塊跑完，否則下一步 embed 會被「正在 reparse」擋下（守門正常運作）
    for _ in range(60):
        d = body(c.get(f"/console/kb/document?id={did}&chunks=0"))
        if not d.get("running") and (d.get("doc") or {}).get("chunk_count", 0) > 0:
            break
        time.sleep(0.25)
    check(not body(c.get(f"/console/kb/document?id={did}&chunks=0")).get("running"),
          "重新分塊完成（背景任務釋放）")

    print("\n[5] embedding → 入庫")
    r = c.post("/console/kb/document/embed", json={"id": did, "mode": "pending"})
    check(r.status_code == 200, "POST /document/embed（背景）", str(body(r)))
    done = False
    for _ in range(120):
        st = c.get(f"/console/kb/document?id={did}&chunks=0")
        d = body(st)
        s = d.get("stats") or {}
        if int(s.get("done") or 0) > 0 and not d.get("running"):
            done = True
            break
        time.sleep(0.5)
    check(done, "向量化完成（輪詢 prog）",
          f"done={((body(c.get(f'/console/kb/document?id={did}&chunks=0')).get('stats')) or {}).get('done')}")

    r = c.post("/console/kb/document/activate", json={"id": did})
    check(r.status_code == 200 and body(r)["doc"]["status"] == "ready", "POST /document/activate")

    # 其餘兩份也入庫（直接走同步函式，省時間）
    import kb_ingest
    for x in doc_ids[1:]:
        kb_ingest.embed_document(x, "pending")
        kb_ingest.activate_document(x)
    ready = kb_db.query("SELECT count(*) c FROM kb_document WHERE collection_id=%s AND status='ready'",
                        (cid,))
    check(int(ready[0]["c"]) == len(doc_ids), f"{len(doc_ids)} 份文件都 ready（{ready[0]['c']}）")

    print("\n[6] 知識塊微調端點")
    ch = kb_db.list_chunks(did)
    if len(ch) >= 2:
        a, b = ch[0], ch[1]
        r = c.post("/console/kb/chunk/update", json={"id": a["id"], "text": "端點手改文字"})
        check(r.status_code == 200 and body(r)["chunk"]["edited"], "POST /chunk/update")

        r = c.post("/console/kb/chunk/split", json={"id": a["id"], "at": 3})
        d = body(r)
        check(r.status_code == 200 and d.get("new_chunk_id"), "POST /chunk/split",
              f"new=#{d.get('new_chunk_id')} 共 {len(d.get('chunks') or [])} 塊")

        r = c.post("/console/kb/chunk/merge", json={"id": a["id"]})
        d = body(r)
        check(r.status_code == 200, "POST /chunk/merge", f"共 {len(d.get('chunks') or [])} 塊")

        r = c.post("/console/kb/chunk/split", json={"id": a["id"], "at": 999999})
        check(r.status_code == 400, "超界拆分位置被擋", body(r).get("error"))

        last = c.get(f"/console/kb/document?id={did}&chunks=0")
        chs = kb_db.list_chunks(did)
        r = c.post("/console/kb/chunk/delete", json={"id": chs[-1]["id"]})
        check(r.status_code == 200, "POST /chunk/delete", f"共 {len(body(r).get('chunks') or [])} 塊")
    else:
        check(False, "知識塊不足，無法測微調")

    print("\n[7] 停用 / 刪除")
    r = c.post("/console/kb/document/enabled", json={"id": did, "enabled": False})
    check(r.status_code == 200 and body(r)["doc"]["enabled"] is False, "POST /document/enabled")
    c.post("/console/kb/document/enabled", json={"id": did, "enabled": True})

    r = c.post("/console/kb/collections/delete", json={"id": cid})
    check(r.status_code == 400, "非空知識庫刪除需二次確認", body(r).get("error", "")[:40])

    # 清理：刪文件（連檔案）+ 集合
    for x in doc_ids:
        c.post("/console/kb/document/delete", json={"id": x})
    left = kb_db.query("SELECT count(*) c FROM kb_document WHERE collection_id=%s", (cid,))
    check(int(left[0]["c"]) == 0, "刪文件後 DB 無殘留")
    r = c.post("/console/kb/collections/delete", json={"id": cid})
    check(r.status_code == 200, "清空後可刪知識庫")

    print("\n[8] 端點齊全性（路由存在）")
    for path in ("/console/kb/overview", "/console/kb/collections", "/console/kb/params",
                 "/console/kb/documents", "/console/kb/doc-types"):
        r = c.get(path)
        check(r.status_code == 200, f"GET {path} → {r.status_code}")

print("\n" + "=" * 72)
print(f" 結果：{'全部通過 ✅' if not FAIL else str(len(FAIL)) + ' 項失敗'}")
for f in FAIL:
    print("   ✗", f)
print("=" * 72)
sys.exit(1 if FAIL else 0)
